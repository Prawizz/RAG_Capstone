"""PPTX extractor: หนึ่งสไลด์ = หนึ่ง Unit (type=slide, method=office_xml)

- ``position = {"slide": N}`` (เริ่มที่ 1 ตามลำดับในไฟล์, สไลด์ซ่อนยังเก็บและใส่ signal ``pptx:hidden`` = 1)
  ``label`` = "สไลด์ N: <หัวเรื่อง>" (หัวเรื่อง = placeholder TITLE / CENTER_TITLE) หรือ "สไลด์ N" ถ้าไม่มี
- **ลำดับการอ่านเรียงตามตำแหน่ง ไม่ใช่ลำดับใน XML** (ลำดับใน XML คือลำดับที่คนสร้างกล่อง ทำให้คู่ข้อความสลับกัน):
  หัวเรื่องก่อน ที่เหลือเรียงจากบนลงล่าง ซ้ายไปขวา กล่องที่ top อยู่ห่างจากกล่องบนสุดของแถวไม่เกิน ``ROW_BAND`` ของ
  ความสูงสไลด์ (5%) นับเป็นแถวเดียวกันแล้วเรียงตาม left; แถวถัดไปเริ่มที่กล่องแรกที่เหลือ
- group shape: ไล่เข้าไปอ่านข้างใน พิกัดของลูกแปลงเป็นพิกัดบนสไลด์ด้วย chOff/chExt ของ group (ซ้อนกันได้)
- Block: หัวเรื่อง = heading level 1, กล่องข้อความ = paragraph (คงการขึ้นบรรทัด) พร้อม bbox หน่วย point,
  ตาราง = table (แถวละบรรทัด ช่องคั่นด้วย " | " เซลล์ผสานไม่ซ้ำ), กราฟ = table ที่ข้อความขึ้นต้น "กราฟ: <ชื่อ>" ตามด้วย
  หมวด | ซีรีส์... อ่านจากข้อมูลในไฟล์ (กราฟที่อ่านไม่ได้ -> signal ``pptx:chart_unreadable`` ไม่ crash),
  speaker notes = paragraph ท้าย unit ขึ้นต้น "บันทึกผู้บรรยาย: "
- ไม่อ่าน placeholder วันที่/ท้ายกระดาษ/เลขสไลด์ ไม่ OCR รูป ไม่อ่าน SmartArt
- signals: char_count, block_count, image_coverage (รูปที่ใหญ่ที่สุด ÷ พื้นที่สไลด์ นิยามเดียวกับ PDF),
  pptx:picture_count และผลนับของ normalizer (norm_fix:* / norm_sig:*)
"""

import hashlib
import time
from dataclasses import dataclass, field
from pathlib import Path
from typing import Callable, Optional

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE, PP_PLACEHOLDER

from app.extraction.errors import ExtractionError
from app.extraction.tally import Tally
from app.models import (
    Block,
    BlockType,
    ErrorCode,
    ExtractionMethod,
    MergedCell,
    QualityFlag,
    QualityReport,
    Table,
    Unit,
    UnitStatus,
    UnitType,
)

ROW_BAND = 0.05
EMU_PER_PT = 12700
NOTES_PREFIX = "บันทึกผู้บรรยาย: "
CHART_PREFIX = "กราฟ: "
_TITLES = (PP_PLACEHOLDER.TITLE, PP_PLACEHOLDER.CENTER_TITLE)
_CHROME = (PP_PLACEHOLDER.DATE, PP_PLACEHOLDER.FOOTER, PP_PLACEHOLDER.SLIDE_NUMBER)

Box = tuple[float, float, float, float]  # x, y, w, h เป็น EMU บนพิกัดสไลด์
Mapper = Callable[[Box], Box]


@dataclass
class PptxResult:
    units: list[Unit] = field(default_factory=list)
    has_outline: bool = False
    producer: Optional[str] = None
    slide_count: int = 0


@dataclass
class _Item:
    kind: str  # text | table | chart | picture
    shape: object
    box: Optional[Box]
    order: int
    is_title: bool = False


def _identity(box: Box) -> Box:
    return box


def _group_mapper(group, parent: Mapper) -> Mapper:
    xf = group._element.xpath("./p:grpSpPr/a:xfrm")
    if not xf:
        return parent
    x = xf[0]
    off, ext, choff, chext = (x.find(f"{{{_A}}}{t}") for t in ("off", "ext", "chOff", "chExt"))
    if None in (off, ext, choff, chext):
        return parent
    ox, oy = int(off.get("x")), int(off.get("y"))
    cx, cy = int(ext.get("cx")), int(ext.get("cy"))
    chx, chy = int(choff.get("x")), int(choff.get("y"))
    chw, chh = int(chext.get("cx")), int(chext.get("cy"))
    sx = cx / chw if chw else 1.0
    sy = cy / chh if chh else 1.0

    def mapper(box: Box) -> Box:
        bx, by, bw, bh = box
        return parent((ox + (bx - chx) * sx, oy + (by - chy) * sy, bw * sx, bh * sy))

    return mapper


_A = "http://schemas.openxmlformats.org/drawingml/2006/main"


def _box(shape) -> Optional[Box]:
    try:
        l, t, w, h = shape.left, shape.top, shape.width, shape.height
    except Exception:
        return None
    if None in (l, t, w, h):
        return None
    return (float(l), float(t), float(w), float(h))


def _shape_type(shape):
    try:
        return shape.shape_type
    except NotImplementedError:
        return None


def _placeholder_type(shape):
    try:
        return shape.placeholder_format.type if shape.is_placeholder else None
    except Exception:
        return None


def _collect(shapes, mapper: Mapper, out: list[_Item]) -> None:
    for sh in shapes:
        st = _shape_type(sh)
        if st == MSO_SHAPE_TYPE.GROUP:
            _collect(sh.shapes, _group_mapper(sh, mapper), out)
            continue
        raw = _box(sh)
        box = mapper(raw) if raw is not None else None
        ph = _placeholder_type(sh)
        if ph in _CHROME:
            continue
        if getattr(sh, "has_chart", False) and sh.has_chart:
            out.append(_Item("chart", sh, box, len(out)))
        elif getattr(sh, "has_table", False) and sh.has_table:
            out.append(_Item("table", sh, box, len(out)))
        elif st == MSO_SHAPE_TYPE.PICTURE or type(sh).__name__ in ("Picture", "PlaceholderPicture"):
            out.append(_Item("picture", sh, box, len(out)))
        elif getattr(sh, "has_text_frame", False) and sh.has_text_frame and sh.text_frame.text.strip():
            out.append(_Item("text", sh, box, len(out), is_title=ph in _TITLES))


def reading_order(items: list[_Item], slide_height: float) -> list[_Item]:
    """เรียงตามแถว (top ห่างจากบนสุดของแถวไม่เกิน ROW_BAND x ความสูงสไลด์) แล้วตาม left; ไม่มีพิกัด = ท้ายสุดตามลำดับเดิม"""
    placed = sorted((i for i in items if i.box is not None), key=lambda i: (i.box[1], i.box[0], i.order))
    unplaced = [i for i in items if i.box is None]
    out: list[_Item] = []
    band = ROW_BAND * slide_height
    k = 0
    while k < len(placed):
        row_top = placed[k].box[1]
        row = []
        while k < len(placed) and placed[k].box[1] - row_top <= band:
            row.append(placed[k])
            k += 1
        out.extend(sorted(row, key=lambda i: (i.box[0], i.box[1], i.order)))
    return out + unplaced


def _bbox(box: Optional[Box]) -> Optional[tuple[float, float, float, float]]:
    if box is None:
        return None
    x, y, w, h = box
    return (x / EMU_PER_PT, y / EMU_PER_PT, (x + w) / EMU_PER_PT, (y + h) / EMU_PER_PT)


def _coverage(box: Optional[Box], slide_w: float, slide_h: float) -> float:
    if box is None or slide_w <= 0 or slide_h <= 0:
        return 0.0
    x, y, w, h = box
    iw = max(0.0, min(x + w, slide_w) - max(x, 0.0))
    ih = max(0.0, min(y + h, slide_h) - max(y, 0.0))
    return iw * ih / (slide_w * slide_h)


def _num(v) -> str:
    if v is None:
        return ""
    if isinstance(v, float):
        return str(int(v)) if v == int(v) and abs(v) < 1e15 else repr(round(v, 6)).rstrip("0").rstrip(".")
    return str(v)


def _table_block(shape, tally: Tally, bbox) -> Optional[Block]:
    tbl = shape.table
    rows: list[list[str]] = []
    merged: list[MergedCell] = []
    lines: list[str] = []
    for r in range(len(tbl.rows)):
        row, line = [], []
        for c in range(len(tbl.columns)):
            cell = tbl.cell(r, c)
            if cell.is_spanned:
                row.append("")
                continue
            text = tally.norm(cell.text_frame.text.replace("\n", " ").replace("\x0b", " "))
            row.append(text)
            if text:
                line.append(text)
            if cell.is_merge_origin:
                merged.append(MergedCell(row=r, col=c, rowspan=cell.span_height, colspan=cell.span_width))
        rows.append(row)
        if line:
            lines.append(" | ".join(line))
    text = "\n".join(lines)
    if not text:
        return None
    header = 1 if getattr(tbl, "first_row", False) else 0
    return Block(type=BlockType.TABLE, text=text, order=0, bbox=bbox,
                 table=Table(rows=rows, header_rows=header, merged_cells=merged))


def _chart_block(shape, tally: Tally, bbox) -> tuple[Optional[Block], bool]:
    """(block, อ่านไม่ได้?)"""
    chart = shape.chart
    title = ""
    try:
        if chart.has_title and chart.chart_title.has_text_frame:
            title = chart.chart_title.text_frame.text
    except Exception:
        pass
    title = tally.norm((title or "").replace("\n", " "))
    head = f"{CHART_PREFIX}{title}" if title else "กราฟ"
    try:
        plots = list(chart.plots)
        cats = ["/".join(lbl) for lbl in plots[0].categories.flattened_labels]
        series = [(s.name, list(s.values)) for p in plots for s in p.series]
        if not cats or not series:
            raise ValueError("no categories or series")
    except Exception:
        block = Block(type=BlockType.PARAGRAPH, text=head, order=0, bbox=bbox) if title else None
        return block, True
    names = [tally.norm(n) or f"ซีรีส์{k}" for k, (n, _) in enumerate(series, 1)]
    grid = [["หมวด", *names]]
    for i, cat in enumerate(cats):
        grid.append([tally.norm(cat), *(_num(v[i]) if i < len(v) else "" for _, v in series)])
    text = "\n".join([head, *(" | ".join(r) for r in grid)])
    return Block(type=BlockType.TABLE, text=text, order=0, bbox=bbox, table=Table(rows=grid, header_rows=1)), False


def _hash(text: str) -> str:
    return hashlib.sha256(text.encode("utf-8")).hexdigest()


def _failed_unit(index: int, n: int, error: Exception, elapsed_ms: float) -> Unit:
    return Unit(
        index=index, type=UnitType.SLIDE, position={"slide": n}, label=f"สไลด์ {n}",
        status=UnitStatus.FAILED,
        reason=f"อ่านสไลด์ {n} ไม่สำเร็จ: {type(error).__name__}: {error}",
        method=ExtractionMethod.OFFICE_XML, text="", text_hash=_hash(""), quality=QualityReport(),
        elapsed_ms=elapsed_ms,
        signals={"char_count": 0.0, "block_count": 0.0, "image_coverage": 0.0, "pptx:picture_count": 0.0},
    )


def _extract_slide(index: int, n: int, slide, slide_w: float, slide_h: float) -> Unit:
    t0 = time.perf_counter()
    tally = Tally()
    items: list[_Item] = []
    _collect(slide.shapes, _identity, items)
    blocks: list[Block] = []
    unreadable = 0

    def add(b: Block) -> None:
        blocks.append(b.model_copy(update={"order": len(blocks)}))

    title_item = next((i for i in items if i.is_title), None)
    title = ""
    if title_item is not None:
        title = tally.norm(title_item.shape.text_frame.text.replace("\n", " ").replace("\x0b", " "))
        if title:
            add(Block(type=BlockType.HEADING, text=title, order=0, level=1, bbox=_bbox(title_item.box)))
    rest = [i for i in items if i is not title_item and i.kind != "picture"]
    for it in reading_order(rest, slide_h):
        bbox = _bbox(it.box)
        if it.kind == "text":
            text = tally.norm(it.shape.text_frame.text)
            if text:
                add(Block(type=BlockType.PARAGRAPH, text=text, order=0, bbox=bbox))
        elif it.kind == "table":
            b = _table_block(it.shape, tally, bbox)
            if b is not None:
                add(b)
        elif it.kind == "chart":
            b, bad = _chart_block(it.shape, tally, bbox)
            unreadable += int(bad)
            if b is not None:
                add(b)
    if slide.has_notes_slide:
        notes = tally.norm(slide.notes_slide.notes_text_frame.text) if slide.notes_slide.notes_text_frame is not None else ""
        if notes:
            add(Block(type=BlockType.PARAGRAPH, text=NOTES_PREFIX + notes, order=0))

    pictures = [i for i in items if i.kind == "picture"]
    coverage = min(max((_coverage(p.box, slide_w, slide_h) for p in pictures), default=0.0), 1.0)
    text = "\n".join(b.text for b in blocks)
    signals: dict[str, float] = {
        "char_count": float(len(text.strip())),
        "block_count": float(len(blocks)),
        "image_coverage": coverage,
        "pptx:picture_count": float(len(pictures)),
        **tally.signal_dict(),
    }
    if slide._element.get("show") == "0":
        signals["pptx:hidden"] = 1.0
    if unreadable:
        signals["pptx:chart_unreadable"] = float(unreadable)
    return Unit(
        index=index,
        type=UnitType.SLIDE,
        position={"slide": n},
        label=f"สไลด์ {n}: {title}" if title else f"สไลด์ {n}",
        status=UnitStatus.OK,
        method=ExtractionMethod.OFFICE_XML,
        blocks=blocks,
        text=text,
        text_hash=_hash(text),
        quality=QualityReport(score=None, flags=[] if text.strip() else [QualityFlag.EMPTY],
                              normalization_changes=sum(tally.fixes.values())),
        elapsed_ms=(time.perf_counter() - t0) * 1000,
        signals=signals,
    )


def extract_pptx(path: str | Path) -> PptxResult:
    try:
        prs = Presentation(str(path))
        slides = list(prs.slides)
    except Exception as e:
        raise ExtractionError(ErrorCode.CORRUPT, f"เปิด PPTX ไม่ได้: {type(e).__name__}: {e}") from e
    slide_w, slide_h = float(prs.slide_width), float(prs.slide_height)
    result = PptxResult(slide_count=len(slides))
    for i, slide in enumerate(slides):
        t0 = time.perf_counter()
        try:
            result.units.append(_extract_slide(i, i + 1, slide, slide_w, slide_h))
        except Exception as e:  # สไลด์เดียวพังไม่ทำให้ทั้งไฟล์พัง
            result.units.append(_failed_unit(i, i + 1, e, (time.perf_counter() - t0) * 1000))
    return result
