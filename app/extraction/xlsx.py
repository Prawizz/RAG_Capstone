"""XLSX extractor: หนึ่งตารางในชีต = หนึ่ง Unit (type=sheet_region, method=office_xml)

ด่านกันไฟล์ใหญ่ (ก่อนอ่านค่าใด ๆ):
1. ขนาดหลังแตกของ ``xl/worksheets/*.xml`` ชีตใดเกิน ``MAX_SHEET_XML_BYTES`` -> too_large (ดูจาก zip ไม่เปิดด้วย openpyxl)
2. ชีตที่มองเห็นมีแถวที่มีค่าจริงเกิน ``MAX_DATA_ROWS`` -> too_large (นับแบบ streaming หยุดทันทีที่เกิน)

หาตาราง (ไม่เชื่อ dimension ของ Excel ดูแต่เซลล์ที่มีค่าจริง):
- ชีต hidden / veryHidden ข้าม นับใน ``xlsx:hidden_sheets`` ของ unit แรกของไฟล์; แถว/คอลัมน์ที่ซ่อนยังอ่าน
- แถวว่างติดกันตั้งแต่ ``BLANK_ROWS_SPLIT`` แถวแบ่งตาราง แถวว่างเดี่ยวไม่แบ่ง
- หัวตาราง = แถวแรกที่ "จำนวนช่องที่มีค่า" >= ครึ่งหนึ่งของจำนวนคอลัมน์ และ >= 2; แถวก่อนหน้า = ชื่อเรื่อง
  เซลล์ผสานแนวนอนนับตามจำนวนคอลัมน์ที่คลุม
  **หัวกลุ่ม vs ชื่อเรื่อง** (สองอย่างหน้าตาเหมือนกัน คือเซลล์เดียวผสานแนวนอน แยกด้วยความกว้างอย่างเดียว):
  ผสานคลุม >= 2 คอลัมน์ แต่*ไม่ถึงทั้งความกว้างตาราง* = หัวกลุ่ม ("ปี 2567" คร่อม 4 จาก 5 คอลัมน์);
  ผสานคลุมทั้งความกว้างตาราง = ชื่อเรื่อง/แบนเนอร์ เหตุผล: ชื่อเรื่องของตารางจริงมักคลุมทั้งหน้า ส่วนหัวกลุ่มต้องเหลือ
  คอลัมน์อื่นไว้ให้หัวอื่น (ไม่งั้นจะไม่มีอะไรให้จัดกลุ่ม) ข้อจำกัด: ชื่อเรื่องที่ผสานแค่บางส่วนของความกว้างจะถูกเข้าใจว่าเป็นหัวกลุ่ม
  หัวหลายชั้น (รวมกันได้): (ก) เซลล์ผสานลงล่างในหัวขยายหัวตารางถึงแถวที่ผสานถึง (ข) แถวหัวที่มีหัวกลุ่มผสานแนวนอนมีแถวถัดไป
  เป็นหัวชั้นล่าง (ค) แถวติดกันเหนือแถวหัวที่มีหัวกลุ่มผสานแนวนอนเป็นหัวชั้นบน (ซ้ำขึ้นไปได้)
  ชื่อคอลัมน์ = ข้อความแต่ละชั้นต่อกันด้วย " / " ไม่ซ้ำ (หัวผสานแนวนอนเติมให้ทุกคอลัมน์ที่คร่อม)
- แถวที่ค่าเหมือนแถวหัวตาราง = หัวซ้ำ ข้าม (``xlsx:repeated_header_rows``)
- แถวมีค่าช่องเดียว: คอลัมน์ 1–2 ของตาราง = ป้ายกลุ่ม "## <ป้าย>" · คอลัมน์อื่น = "หมายเหตุ: <ข้อความ>"
- แถวข้อมูล "ชื่อคอลัมน์: ค่า | ..." ข้ามช่องว่างแต่เก็บ "-"; เซลล์ผสานแนวตั้งในข้อมูลเติมค่าให้ทุกแถวที่คลุม
- ตารางที่ไม่มีหัวตาราง: ชื่อคอลัมน์ = ตัวอักษรคอลัมน์ของ Excel, แถวค่าเดียวเป็น paragraph ธรรมดา
- ไม่ทายชั้นของป้ายกลุ่ม: ต้น unit ใส่ "[ป้ายก่อนหน้า: ก → ข → ค]" (ป้าย ``CONTEXT_LABELS`` อันล่าสุดตามลำดับที่เจอ)
  signal ``xlsx:group_label_max_repeat`` (ป้ายที่ซ้ำมากสุดในตารางนั้น) เก็บไว้ดูเท่านั้น ห้ามใช้ตัดสิน

ค่าในเซลล์: ค่าที่ Excel คำนวณเก็บไว้ (data_only) สูตรที่ไม่มีค่าเก็บไว้ -> ว่าง + ``xlsx:formula_no_cache``;
จำนวนเต็มไม่มี ".0"; รูปแบบ % -> x100 + "%"; วันที่ YYYY-MM-DD (มีเวลาเฉพาะเมื่อไม่ใช่ 00:00); error เป็นข้อความ

Unit: ข้อความไม่เกิน ``MAX_UNIT_CHARS`` (นับชื่อเรื่อง ไม่นับบรรทัดบริบท) เกินแบ่งที่ขอบแถว (ป้ายท้าย unit ย้ายไป unit ถัดไป)
ข้อความ unit = ชื่อเรื่อง -> [ป้ายก่อนหน้า] -> แถว; ``position = {"sheet", "range"}``
label = "ชีต <ชื่อ> แถว a–b[: <ชื่อเรื่อง>]"
Block: heading (ชื่อเรื่อง) · paragraph (บรรทัดบริบท, "## ป้าย", หมายเหตุ) · table (หัวตาราง + แถวข้อมูลที่ต่อเนื่องกัน)
signals: char_count, block_count, xlsx:data_rows, xlsx:merged_cells, xlsx:hidden_sheets, xlsx:repeated_header_rows,
xlsx:group_label_max_repeat, xlsx:formula_no_cache และผลนับของ normalizer
"""

import hashlib
import re
import time
import zipfile
from collections import Counter
from dataclasses import dataclass, field
from datetime import date, datetime, time as dtime
from pathlib import Path
from typing import Optional

import openpyxl
from openpyxl.utils import get_column_letter

from app.extraction.errors import ExtractionError
from app.extraction.tally import Tally
from app.models import (
    Block,
    BlockType,
    ErrorCode,
    ExtractionMethod,
    MergedCell,
    QualityReport,
    Table,
    Unit,
    UnitStatus,
    UnitType,
)

MAX_SHEET_XML_BYTES = 50 * 1024 * 1024
MAX_DATA_ROWS = 50_000
MAX_UNIT_CHARS = 4000
BLANK_ROWS_SPLIT = 2
CONTEXT_LABELS = 3
GROUP_LABEL_MAX_COL = 2
TOO_LARGE_REASON = "ไฟล์เป็นฐานข้อมูลขนาดใหญ่ ไม่เหมาะกับการค้นแบบเอกสาร"

_SHEET_XML = re.compile(r"^xl/worksheets/[^/]+\.xml$")


@dataclass
class XlsxResult:
    units: list[Unit]
    has_outline: bool = False
    producer: Optional[str] = None
    page_equivalent: int = 0


@dataclass
class _Sheet:
    name: str
    cells: dict  # (row, col) -> ข้อความที่จัดรูปแล้ว (เฉพาะเซลล์ที่มีค่า)
    merges: list  # (r1, c1, r2, c2)
    no_cache: set  # (row, col) ของสูตรที่ไม่มีค่าเก็บไว้
    norm: dict = field(default_factory=dict)  # (row, col) -> Tally ของเซลล์ที่ normalizer แก้อะไรไว้


@dataclass
class _Item:
    kind: str  # data | label | note | text | repeat
    row: int
    text: str = ""
    cells: list = field(default_factory=list)  # kind=data: ค่าต่อคอลัมน์
    ctx: tuple = ()  # ป้าย CONTEXT_LABELS อันล่าสุดก่อน item นี้


def _hash(text: str) -> str:
    return hashlib.sha256(text.encode("utf-8")).hexdigest()


def _too_large(detail: str) -> ExtractionError:
    return ExtractionError(ErrorCode.TOO_LARGE, f"{TOO_LARGE_REASON} ({detail})")


# ---- ด่านกันไฟล์ใหญ่ -------------------------------------------------------------------------------


def check_sheet_sizes(path: Path) -> None:
    """ด่านแรก: ดูขนาดหลังแตกใน zip โดยไม่เปิดไฟล์ด้วย openpyxl"""
    try:
        with zipfile.ZipFile(path) as z:
            for info in z.infolist():
                if _SHEET_XML.match(info.filename) and info.file_size > MAX_SHEET_XML_BYTES:
                    raise _too_large(
                        f"{info.filename} หลังแตกไฟล์ {info.file_size / 1024 / 1024:.0f} MB "
                        f"เกิน {MAX_SHEET_XML_BYTES // 1024 // 1024} MB"
                    )
    except zipfile.BadZipFile as e:
        raise ExtractionError(ErrorCode.CORRUPT, f"ไฟล์ ZIP เสียหาย: {e}") from e


def check_row_counts(path: Path) -> None:
    """ด่านที่สอง: นับแถวที่มีค่าจริงของชีตที่มองเห็นแบบ streaming หยุดทันทีที่เกิน"""
    try:
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    except Exception as e:
        raise ExtractionError(ErrorCode.CORRUPT, f"เปิดสมุดงานไม่ได้: {type(e).__name__}: {e}") from e
    try:
        for ws in wb.worksheets:
            if ws.sheet_state != "visible":
                continue
            rows = 0
            for row in ws.iter_rows(values_only=True):
                if any(v is not None and str(v).strip() != "" for v in row):
                    rows += 1
                    if rows > MAX_DATA_ROWS:
                        raise _too_large(f"ชีต {ws.title} มีแถวข้อมูลเกิน {MAX_DATA_ROWS:,} แถว")
    finally:
        wb.close()


# ---- ค่าในเซลล์ -------------------------------------------------------------------------------------


def _num_text(v) -> str:
    if isinstance(v, int):
        return str(v)
    f = float(v)
    if f == int(f) and abs(f) < 1e15:
        return str(int(f))
    return f"{f:.15g}"


def _format_value(cell, tally: Tally) -> str:
    v = cell.value
    if v is None:
        return ""
    if isinstance(v, bool):
        return "TRUE" if v else "FALSE"
    if isinstance(v, datetime):
        if v.hour == v.minute == v.second == v.microsecond == 0:
            return v.strftime("%Y-%m-%d")
        return v.strftime("%Y-%m-%d %H:%M:%S" if v.second else "%Y-%m-%d %H:%M")
    if isinstance(v, date):
        return v.strftime("%Y-%m-%d")
    if isinstance(v, dtime):
        return v.strftime("%H:%M:%S" if v.second else "%H:%M")
    if isinstance(v, (int, float)) or type(v).__name__ == "Decimal":
        fmt = cell.number_format or ""
        if "%" in fmt:
            return _num_text(round(float(v) * 100, 10)) + "%"
        if isinstance(v, int) and v >= 0 and re.fullmatch(r"0{2,}", fmt):
            return str(v).zfill(len(fmt))  # เลขศูนย์นำหน้าตามรูปแบบเซลล์
        return _num_text(v)
    return tally.norm(str(v).strip())


def _has_formulas(path: Path) -> bool:
    with zipfile.ZipFile(path) as z:
        for info in z.infolist():
            if _SHEET_XML.match(info.filename):
                data = z.read(info.filename)
                if b"<f>" in data or b"<f " in data:
                    return True
    return False


def _load_sheets(path: Path) -> tuple[list[_Sheet], int]:
    try:
        wb = openpyxl.load_workbook(path, data_only=True)
        wb_f = openpyxl.load_workbook(path, data_only=False) if _has_formulas(path) else None
    except ExtractionError:
        raise
    except Exception as e:
        raise ExtractionError(ErrorCode.CORRUPT, f"เปิดสมุดงานไม่ได้: {type(e).__name__}: {e}") from e
    sheets, hidden = [], 0
    for ws in wb.worksheets:
        if ws.sheet_state != "visible":
            hidden += 1
            continue
        formula_ws = wb_f[ws.title] if wb_f is not None else None
        cells, no_cache, norm = {}, set(), {}
        for (r, c), cell in ws._cells.items():  # ไม่ใช้ iter_rows: dimension ที่ประกาศอาจใหญ่เกินจริง
            t = Tally()
            text = _format_value(cell, t)
            if text:
                cells[(r, c)] = text
                if t.fixes or t.signals:
                    norm[(r, c)] = t
            elif cell.value is None and formula_ws is not None:
                fc = formula_ws._cells.get((r, c))
                if fc is not None and fc.data_type == "f":
                    no_cache.add((r, c))
        merges = [(m.min_row, m.min_col, m.max_row, m.max_col) for m in ws.merged_cells.ranges]
        sheets.append(_Sheet(ws.title, cells, merges, no_cache, norm))
    return sheets, hidden


# ---- หาตาราง ---------------------------------------------------------------------------------------


def _vertical_fill(sheet: _Sheet) -> dict:
    """ค่าของเซลล์ + ค่าของเซลล์ผสานแนวตั้งที่เติมลงแถวที่คลุม (เฉพาะคอลัมน์ต้นทาง)"""
    filled = dict(sheet.cells)
    last_row = max((r for r, _ in sheet.cells), default=0)
    for r1, c1, r2, _c2 in sheet.merges:
        if r2 > r1 and (r1, c1) in sheet.cells:
            for r in range(r1 + 1, min(r2, last_row) + 1):
                filled.setdefault((r, c1), sheet.cells[(r1, c1)])
    return filled


def _split_tables(rows: list[int]) -> list[list[int]]:
    tables, cur = [], [rows[0]]
    for r in rows[1:]:
        if r - cur[-1] - 1 >= BLANK_ROWS_SPLIT:
            tables.append(cur)
            cur = []
        cur.append(r)
    tables.append(cur)
    return tables


def _merged_value(sheet: _Sheet, r: int, c: int) -> str:
    v = sheet.cells.get((r, c))
    if v:
        return v
    for r1, c1, r2, c2 in sheet.merges:
        if r1 <= r <= r2 and c1 <= c <= c2:
            return sheet.cells.get((r1, c1), "")
    return ""


def _flat(text: str) -> str:
    return re.sub(r"\s*[\r\n]+\s*", " ", text)


def _build_items(sheet: _Sheet, filled: dict, rs: list[int], cols: list[int]):
    """คืน (title_rows, items, header_rows, names, header_end_row, repeated)"""
    mincol = cols[0]
    ncols = len(cols)
    need = max(2, (ncols + 1) // 2)
    row_vals = {r: {c: filled[(r, c)] for c in cols if (r, c) in filled} for r in rs}
    rowset = set(rs)

    def group_spans(r: int) -> list[int]:
        """ผสานแนวนอนที่เริ่มในแถว r และคลุม >= 2 คอลัมน์ แต่ไม่ถึงทั้งความกว้างตาราง (= หัวกลุ่ม) -> จำนวนคอลัมน์ที่คลุม"""
        out = []
        for r1, c1, _r2, c2 in sheet.merges:
            covered = sum(1 for c in cols if c1 <= c <= c2)
            if r1 == r and c1 in row_vals.get(r, {}) and 2 <= covered < ncols:
                out.append(covered)
        return out

    def filled_count(r: int) -> int:
        return len(row_vals[r]) + sum(n - 1 for n in group_spans(r))

    h0 = next((r for r in rs if filled_count(r) >= need), None)

    names = {c: get_column_letter(c) for c in cols}
    title: list[str] = []
    header = end = None
    if h0 is not None:
        header = h0
        i = rs.index(h0)
        while i > 0 and rs[i - 1] == header - 1 and group_spans(rs[i - 1]):  # หัวชั้นบนที่เป็นหัวกลุ่มผสานแนวนอน
            header = rs[i - 1]
            i -= 1
        end = h0
        grown = True
        while grown:
            grown = False
            for r1, c1, r2, _c2 in sheet.merges:  # เซลล์ผสานลงล่างขยายหัวตารางถึงแถวที่ผสานถึง
                if header <= r1 <= end and c1 in cols and r2 > end:
                    end, grown = r2, True
            if end + 1 in rowset and group_spans(end):  # หัวกลุ่มแนวนอนมีหัวชั้นล่างต่อ
                end, grown = end + 1, True
        end = min(end, rs[-1])
        for c in cols:
            parts: list[str] = []
            for r in range(header, end + 1):
                v = _merged_value(sheet, r, c)
                if v and v not in parts:
                    parts.append(v)
            if parts:
                names[c] = " / ".join(parts)
        title = [_flat(" ".join(row_vals[r][c] for c in cols if c in row_vals[r])) for r in rs if r < header]
    body = [r for r in rs if header is None or r > end]

    items: list[_Item] = []
    history: list[str] = []
    repeated = 0
    skip_until = -1
    for r in body:
        if r <= skip_until:
            continue
        vals = row_vals[r]
        ctx = tuple(history[-CONTEXT_LABELS:])
        if header is not None and vals == row_vals[header]:
            repeated += 1
            skip_until = r + (end - header)
            items.append(_Item("repeat", r, ctx=ctx))
        elif len(vals) == 1:
            (c, v), = vals.items()
            v = _flat(v)
            if header is None:
                items.append(_Item("text", r, v, ctx=ctx))
            elif c - mincol + 1 <= GROUP_LABEL_MAX_COL:
                items.append(_Item("label", r, f"## {v}", ctx=ctx))
                if not history or history[-1] != v:  # ยุบป้ายที่ซ้ำติดกัน
                    history.append(v)
            else:
                items.append(_Item("note", r, f"หมายเหตุ: {v}", ctx=ctx))
        else:
            line = " | ".join(f"{names[c]}: {_flat(vals[c])}" for c in cols if c in vals)
            items.append(_Item("data", r, line, [vals.get(c, "") for c in cols], ctx))
    return title, items, header, names, end, repeated


def _pack(items: list[_Item], base: int) -> list[list[_Item]]:
    """แบ่ง item เป็นกลุ่มที่ข้อความ (รวมชื่อเรื่อง ``base`` ตัว ไม่รวมบรรทัดบริบท) ไม่เกิน MAX_UNIT_CHARS ที่ขอบแถว
    ป้ายท้ายกลุ่มย้ายไปกลุ่มถัดไป"""
    groups: list[list[_Item]] = []
    cur: list[_Item] = []
    size = base
    for it in items:
        n = len(it.text) + 1 if it.text else 0
        if cur and size + n > MAX_UNIT_CHARS:
            k = len(cur)
            while k > 0 and cur[k - 1].kind == "label":
                k -= 1
            if k > 0:
                groups.append(cur[:k])
                cur = cur[k:]
                size = base + sum(len(x.text) + 1 for x in cur if x.text)
        cur.append(it)
        size += n
    if cur:
        groups.append(cur)
    return groups


def _table_blocks(
    group: list[_Item], names_row: list[str], cols: list[int], merges, start_order: int
) -> list[Block]:
    blocks: list[Block] = []
    order = start_order
    i = 0
    while i < len(group):
        it = group[i]
        if it.kind == "data":
            run = []
            while i < len(group) and group[i].kind in ("data", "repeat"):
                if group[i].kind == "data":
                    run.append(group[i])
                i += 1
            run_rows = [x.row for x in run]
            merged = []
            for r1, c1, r2, c2 in merges:
                if r1 in run_rows and c1 in cols and (r2 > r1 or c2 > c1):
                    covered = [c for c in cols if c1 <= c <= c2]
                    rowspan = sum(1 for r in run_rows if r1 <= r <= r2)
                    merged.append(
                        MergedCell(row=run_rows.index(r1) + 1, col=cols.index(c1), rowspan=max(rowspan, 1),
                                   colspan=max(len(covered), 1))
                    )
            text = "\n".join(x.text for x in run)
            blocks.append(Block(
                type=BlockType.TABLE, text=text, order=order,
                table=Table(rows=[names_row] + [x.cells for x in run], header_rows=1, merged_cells=merged),
            ))
            order += 1
            continue
        if it.kind != "repeat":
            blocks.append(Block(type=BlockType.PARAGRAPH, text=it.text, order=order))
            order += 1
        i += 1
    return blocks


def _extract_table(
    sheet: _Sheet, filled: dict, rs: list[int], index0: int
) -> list[Unit]:
    t0 = time.perf_counter()
    rowset = set(rs)
    cols = sorted({c for (r, c) in filled if r in rowset})
    title, items, header, names, _end, repeated = _build_items(sheet, filled, rs, cols)
    names_row = [names[c] for c in cols]
    first_row, last_row = rs[0], rs[-1]
    label_counts = Counter(x.text[3:] for x in items if x.kind == "label")
    max_repeat = max(label_counts.values(), default=0)
    title_text = "\n".join(title)
    groups = _pack(items, len(title_text) + 1 if title_text else 0) or [[]]
    rng_cols = f"{get_column_letter(cols[0])}", f"{get_column_letter(cols[-1])}"
    units: list[Unit] = []
    for gi, group in enumerate(groups):
        blocks: list[Block] = []
        order = 0
        if title_text:
            blocks.append(Block(type=BlockType.HEADING, text=title_text, order=order, level=1))
            order += 1
        if group and group[0].ctx:
            blocks.append(Block(
                type=BlockType.PARAGRAPH, order=order,
                text="[ป้ายก่อนหน้า: " + " → ".join(group[0].ctx) + "]",
            ))
            order += 1
        blocks.extend(_table_blocks(group, names_row, cols, sheet.merges, order))
        for k, b in enumerate(blocks):
            b.order = k
        text = "\n".join(b.text for b in blocks)
        a = first_row if gi == 0 else group[0].row
        b_ = last_row if gi == len(groups) - 1 else group[-1].row
        in_unit = lambda r: a <= r <= b_  # noqa: E731
        in_cols = lambda c: cols[0] <= c <= cols[-1]  # noqa: E731
        signals = {
            "char_count": float(len(text)),
            "block_count": float(len(blocks)),
            "xlsx:data_rows": float(sum(1 for x in group if x.kind == "data")),
            "xlsx:merged_cells": float(sum(
                1 for r1, c1, r2, c2 in sheet.merges if in_unit(r1) and in_cols(c1) and (r2 > r1 or c2 > c1)
            )),
        }
        rep = sum(1 for x in group if x.kind == "repeat")
        if rep:
            signals["xlsx:repeated_header_rows"] = float(rep)
        if max_repeat:
            signals["xlsx:group_label_max_repeat"] = float(max_repeat)
        missing = sum(1 for (r, c) in sheet.no_cache if in_unit(r) and in_cols(c))
        if missing:
            signals["xlsx:formula_no_cache"] = float(missing)
        tally = Tally()
        for (r, c), t in sheet.norm.items():
            if in_unit(r) and in_cols(c):
                tally.add(t)
        signals.update(tally.signal_dict())
        label = f"ชีต {sheet.name} แถว {a}–{b_}" + (f": {_flat(title_text)}" if title_text else "")
        units.append(Unit(
            index=index0 + gi,
            type=UnitType.SHEET_REGION,
            position={"sheet": sheet.name, "range": f"{rng_cols[0]}{a}:{rng_cols[1]}{b_}"},
            label=label,
            status=UnitStatus.OK,
            method=ExtractionMethod.OFFICE_XML,
            blocks=blocks,
            text=text,
            text_hash=_hash(text),
            quality=QualityReport(score=None, normalization_changes=sum(tally.fixes.values())),
            signals=signals,
        ))
    elapsed = (time.perf_counter() - t0) * 1000 / len(units)
    for u in units:
        u.elapsed_ms = elapsed
    return units


def extract_xlsx(path: str | Path) -> XlsxResult:
    path = Path(path)
    check_sheet_sizes(path)
    check_row_counts(path)
    sheets, hidden = _load_sheets(path)
    units: list[Unit] = []
    for sheet in sheets:
        if not sheet.cells:
            continue
        filled = _vertical_fill(sheet)
        for rs in _split_tables(sorted({r for r, _ in filled})):
            units.extend(_extract_table(sheet, filled, rs, len(units)))
    if units and hidden:
        units[0].signals["xlsx:hidden_sheets"] = float(hidden)
    return XlsxResult(units=units, page_equivalent=len(units))
