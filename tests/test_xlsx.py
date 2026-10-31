"""XLSX extractor (app/extraction/xlsx.py) ด้วยไฟล์สังเคราะห์จาก openpyxl"""

import datetime as dt
import zipfile

import openpyxl
from openpyxl.styles import PatternFill
from openpyxl.workbook import Workbook

from app.extraction import xlsx as xlsx_mod
from app.extraction.pipeline import extract_file
from app.models import BlockType, DetectedType, DocumentStatus, ErrorCode, UnitType

HEAD = ["รหัส", "ชื่อ", "เบอร์ต่อ"]


def save(wb, tmp_path, name="a.xlsx"):
    f = tmp_path / name
    wb.save(f)
    return f


def extract(path):
    d = extract_file(path)
    assert d.error is None, d.error
    return d


def lines(unit):
    return unit.text.split("\n")


def simple_sheet(ws, header=HEAD, rows=((1, "สมชาย", 1234), (2, "สมหญิง", 5678)), top=1):
    for j, h in enumerate(header, 1):
        ws.cell(top, j, h)
    for i, r in enumerate(rows, top + 1):
        for j, v in enumerate(r, 1):
            ws.cell(i, j, v)


# ---------- พื้นฐาน / pipeline ----------

def test_basic_table_row_format_and_contract(tmp_path):
    wb = Workbook()
    ws = wb.active
    ws.title = "เบอร์ต่อ"
    simple_sheet(ws)
    d = extract(save(wb, tmp_path))
    assert d.detected_type == DetectedType.XLSX and d.status == DocumentStatus.DONE
    assert d.page_equivalent == 1
    (u,) = d.units
    assert u.type == UnitType.SHEET_REGION and u.method.value == "office_xml"
    assert u.position == {"sheet": "เบอร์ต่อ", "range": "A1:C3"}
    assert u.label == "ชีต เบอร์ต่อ แถว 1–3"
    assert lines(u) == [
        "รหัส: 1 | ชื่อ: สมชาย | เบอร์ต่อ: 1234",
        "รหัส: 2 | ชื่อ: สมหญิง | เบอร์ต่อ: 5678",
    ]
    (b,) = u.blocks
    assert b.type == BlockType.TABLE
    assert b.table.rows == [HEAD, ["1", "สมชาย", "1234"], ["2", "สมหญิง", "5678"]]
    assert b.table.header_rows == 1
    assert u.signals["xlsx:data_rows"] == 2 and u.signals["block_count"] == 1
    assert u.signals["char_count"] == len(u.text)


def test_multiple_sheets_in_order(tmp_path):
    wb = Workbook()
    simple_sheet(wb.active)
    wb.active.title = "ก"
    simple_sheet(wb.create_sheet("ข"))
    d = extract(save(wb, tmp_path))
    assert [u.position["sheet"] for u in d.units] == ["ก", "ข"]
    assert [u.index for u in d.units] == [0, 1] and d.page_equivalent == 2


def test_hidden_sheets_skipped_and_counted_on_first_unit(tmp_path):
    wb = Workbook()
    simple_sheet(wb.active)
    wb.active.title = "เห็น"
    h = wb.create_sheet("ซ่อน")
    simple_sheet(h)
    h.sheet_state = "hidden"
    v = wb.create_sheet("ซ่อนมาก")
    simple_sheet(v)
    v.sheet_state = "veryHidden"
    simple_sheet(wb.create_sheet("เห็น2"))
    d = extract(save(wb, tmp_path))
    assert [u.position["sheet"] for u in d.units] == ["เห็น", "เห็น2"]
    assert d.units[0].signals["xlsx:hidden_sheets"] == 2
    assert "xlsx:hidden_sheets" not in d.units[1].signals


def test_hidden_rows_and_columns_are_still_read(tmp_path):
    wb = Workbook()
    ws = wb.active
    simple_sheet(ws)
    ws.row_dimensions[2].hidden = True
    ws.column_dimensions["B"].hidden = True
    (u,) = extract(save(wb, tmp_path)).units
    assert "รหัส: 1 | ชื่อ: สมชาย" in u.text


# ---------- ขอบเขตและการแบ่งตาราง ----------

def test_two_blank_rows_split_one_blank_row_does_not(tmp_path):
    wb = Workbook()
    ws = wb.active
    simple_sheet(ws, top=1)  # แถว 1-3
    simple_sheet(ws, header=["สินค้า", "ราคา"], rows=[("น้ำ", 10)], top=6)  # ว่าง 4,5 -> ตารางที่ 2
    simple_sheet(ws, header=["ก", "ข"], rows=[("x", 1)], top=9)  # ว่างแถว 8 เดียว -> ต่อกับตารางที่ 2
    d = extract(save(wb, tmp_path))
    assert [u.position["range"] for u in d.units] == ["A1:C3", "A6:B10"]
    assert "สินค้า: น้ำ | ราคา: 10" in d.units[1].text


def test_single_blank_row_inside_table_keeps_one_unit(tmp_path):
    wb = Workbook()
    ws = wb.active
    simple_sheet(ws, rows=[(1, "ก", 1), (2, "ข", 2)])
    ws.cell(5, 1, 3), ws.cell(5, 2, "ค"), ws.cell(5, 3, 3)  # ว่างแถว 4 เดียว
    (u,) = extract(save(wb, tmp_path)).units
    assert u.signals["xlsx:data_rows"] == 3


def test_bounds_come_from_values_not_declared_dimension(tmp_path):
    wb = Workbook()
    ws = wb.active
    simple_sheet(ws)
    ws["Z1000"].fill = PatternFill("solid", fgColor="FFFF00")  # มี style แต่ไม่มีค่า
    f = save(wb, tmp_path)
    assert openpyxl.load_workbook(f).active.dimensions == "A1:Z1000"
    (u,) = extract(f).units
    assert u.position["range"] == "A1:C3"


def test_workbook_with_all_sheets_empty_or_hidden_is_empty_error(tmp_path):
    wb = Workbook()
    wb.active.title = "ว่าง"
    h = wb.create_sheet("ซ่อน")
    simple_sheet(h)
    h.sheet_state = "hidden"
    d = extract_file(save(wb, tmp_path))
    assert d.status == DocumentStatus.FAILED and d.unit_count == 0
    assert d.error.code == ErrorCode.EMPTY and d.error.message == "ไฟล์ไม่มีเนื้อหาที่อ่านได้" and not d.error.retriable


def test_empty_sheet_gives_no_unit(tmp_path):
    wb = Workbook()
    simple_sheet(wb.active)
    wb.create_sheet("ว่าง")
    d = extract(save(wb, tmp_path))
    assert len(d.units) == 1


# ---------- หัวตาราง ----------

def test_title_rows_before_header_and_label(tmp_path):
    wb = Workbook()
    ws = wb.active
    ws.title = "S"
    ws["A1"] = "รายชื่อพนักงาน"
    ws["A2"] = "ข้อมูล ณ ต้นปี"
    simple_sheet(ws, top=4)
    (u,) = extract(save(wb, tmp_path)).units
    assert lines(u)[:2] == ["รายชื่อพนักงาน", "ข้อมูล ณ ต้นปี"]
    assert u.blocks[0].type == BlockType.HEADING and u.blocks[0].level == 1
    assert u.label == "ชีต S แถว 1–6: รายชื่อพนักงาน ข้อมูล ณ ต้นปี"
    assert u.position["range"] == "A1:C6"


def test_header_needs_half_of_columns_and_two_cells(tmp_path):
    wb = Workbook()
    ws = wb.active
    ws["A1"] = "หัวเรื่อง"
    ws["A2"] = "x"
    ws["B2"] = "y"  # 2 ช่องจาก 4 คอลัมน์ = ครึ่งพอดี -> หัวตาราง
    ws["A3"], ws["B3"], ws["C3"], ws["D3"] = 1, 2, 3, 4
    (u,) = extract(save(wb, tmp_path)).units
    assert lines(u)[0] == "หัวเรื่อง"
    assert "x: 1 | y: 2 | C: 3 | D: 4" in u.text  # คอลัมน์ที่ไม่มีชื่อใช้ตัวอักษรของ Excel


def test_two_level_header_merged_down_and_across(tmp_path):
    wb = Workbook()
    ws = wb.active
    for col, name in zip("ABC", ["รหัส", "ชื่อ", "ภาค"]):
        ws[f"{col}1"] = name
        ws.merge_cells(f"{col}1:{col}2")
    ws["D1"] = "ปี 2567"
    ws.merge_cells("D1:G1")
    for col, q in zip("DEFG", ["Q1", "Q2", "Q3", "Q4"]):
        ws[f"{col}2"] = q
    for j, v in enumerate([1, "ก", "เหนือ", 10, 20, 30, 40], 1):
        ws.cell(3, j, v)
    (u,) = extract(save(wb, tmp_path)).units
    (b,) = u.blocks
    assert b.table.rows[0] == ["รหัส", "ชื่อ", "ภาค", "ปี 2567 / Q1", "ปี 2567 / Q2", "ปี 2567 / Q3", "ปี 2567 / Q4"]
    assert lines(u) == [
        "รหัส: 1 | ชื่อ: ก | ภาค: เหนือ | ปี 2567 / Q1: 10 | ปี 2567 / Q2: 20 | ปี 2567 / Q3: 30 | ปี 2567 / Q4: 40"
    ]
    assert u.position["range"] == "A1:G3"


def put_row(ws, r, values, start=1):
    for j, v in enumerate(values, start):
        ws.cell(r, j, v)


def test_two_level_header_horizontal_merge_only(tmp_path):
    wb = Workbook()
    ws = wb.active
    ws["B1"] = "ปี 2567"
    ws.merge_cells("B1:E1")
    put_row(ws, 2, ["ภาค", "ไตรมาส 1", "ไตรมาส 2", "ไตรมาส 3", "ไตรมาส 4"])
    put_row(ws, 3, ["เหนือ", 1, 2, 3, 4])
    (u,) = extract(save(wb, tmp_path)).units
    assert u.blocks[0].table.rows[0] == [
        "ภาค", "ปี 2567 / ไตรมาส 1", "ปี 2567 / ไตรมาส 2", "ปี 2567 / ไตรมาส 3", "ปี 2567 / ไตรมาส 4"]
    assert lines(u) == ["ภาค: เหนือ | ปี 2567 / ไตรมาส 1: 1 | ปี 2567 / ไตรมาส 2: 2 | ปี 2567 / ไตรมาส 3: 3 | ปี 2567 / ไตรมาส 4: 4"]
    assert u.position["range"] == "A1:E3"


def test_two_level_header_small_group_found_from_lower_row(tmp_path):
    # หัวกลุ่มคลุมแค่ 2 จาก 7 คอลัมน์ (นับได้ 2 < ครึ่ง) แต่แถวล่างเต็ม -> แถวบนต้องเป็นหัวชั้นบน ไม่ใช่ชื่อเรื่อง
    wb = Workbook()
    ws = wb.active
    ws["C1"] = "ปี 2567"
    ws.merge_cells("C1:D1")
    put_row(ws, 2, list("กขคงจฉช"))
    put_row(ws, 3, [1, 2, 3, 4, 5, 6, 7])
    (u,) = extract(save(wb, tmp_path)).units
    assert u.blocks[0].table.rows[0] == ["ก", "ข", "ปี 2567 / ค", "ปี 2567 / ง", "จ", "ฉ", "ช"]
    assert not any(b.type == BlockType.HEADING for b in u.blocks)


def test_two_level_header_mixed_vertical_and_horizontal(tmp_path):
    wb = Workbook()
    ws = wb.active
    ws["A1"] = "จังหวัด"
    ws.merge_cells("A1:A2")
    ws["B1"] = "ปี 2567"
    ws.merge_cells("B1:C1")
    ws["D1"] = "รวม"
    ws.merge_cells("D1:D2")
    put_row(ws, 2, ["ไตรมาส 1", "ไตรมาส 2"], start=2)
    put_row(ws, 3, ["เชียงใหม่", 10, 20, 30])
    (u,) = extract(save(wb, tmp_path)).units
    assert u.blocks[0].table.rows[0] == ["จังหวัด", "ปี 2567 / ไตรมาส 1", "ปี 2567 / ไตรมาส 2", "รวม"]
    assert lines(u) == ["จังหวัด: เชียงใหม่ | ปี 2567 / ไตรมาส 1: 10 | ปี 2567 / ไตรมาส 2: 20 | รวม: 30"]


def test_full_width_merged_title_stays_title_above_group_header(tmp_path):
    wb = Workbook()
    ws = wb.active
    ws["A1"] = "สถิติรายไตรมาส"
    ws.merge_cells("A1:E1")  # คลุมทั้งความกว้างตาราง = ชื่อเรื่อง
    ws["B2"] = "ปี 2567"
    ws.merge_cells("B2:E2")
    put_row(ws, 3, ["ภาค", "ไตรมาส 1", "ไตรมาส 2", "ไตรมาส 3", "ไตรมาส 4"])
    put_row(ws, 4, ["เหนือ", 1, 2, 3, 4])
    (u,) = extract(save(wb, tmp_path)).units
    assert u.blocks[0].type == BlockType.HEADING and u.blocks[0].text == "สถิติรายไตรมาส"
    table = [b for b in u.blocks if b.type == BlockType.TABLE][0]
    assert table.table.rows[0][0] == "ภาค" and table.table.rows[0][1] == "ปี 2567 / ไตรมาส 1"
    assert u.label.endswith(": สถิติรายไตรมาส")


def test_full_width_merged_row_alone_is_title_not_header(tmp_path):
    wb = Workbook()
    ws = wb.active
    ws["A1"] = "แบนเนอร์"
    ws.merge_cells("A1:C1")
    simple_sheet(ws, top=2)
    (u,) = extract(save(wb, tmp_path)).units
    assert u.blocks[0].type == BlockType.HEADING and u.blocks[1].table.rows[0] == HEAD


def test_repeated_group_header_is_skipped_with_its_sub_row(tmp_path):
    wb = Workbook()
    ws = wb.active
    for top in (1, 5):
        ws.cell(top, 2, "ปี 2567")
        ws.merge_cells(start_row=top, start_column=2, end_row=top, end_column=3)
        put_row(ws, top + 1, ["ภาค", "ไตรมาส 1", "ไตรมาส 2"])
    put_row(ws, 3, ["เหนือ", 1, 2])
    put_row(ws, 7, ["ใต้", 3, 4])
    (u,) = extract(save(wb, tmp_path)).units
    assert u.signals["xlsx:repeated_header_rows"] == 1 and u.signals["xlsx:data_rows"] == 2


def test_context_collapses_adjacent_duplicate_labels(tmp_path, monkeypatch):
    monkeypatch.setattr(xlsx_mod, "MAX_UNIT_CHARS", 160)
    wb = Workbook()
    ws = wb.active
    put_row(ws, 1, HEAD)
    r = 2
    for g in ["ก", "ข", "ข", "ข", "ค", "ค"]:
        ws.cell(r, 2, g)
        r += 1
    for k in range(1, 8):
        put_row(ws, r, [k, "คนที่ " + "x" * 30, k])
        r += 1
    d = extract(save(wb, tmp_path))
    ctx = [ln for ln in lines(d.units[1]) if ln.startswith("[ป้ายก่อนหน้า")]
    assert ctx == ["[ป้ายก่อนหน้า: ก → ข → ค]"]
    assert d.units[0].signals["xlsx:group_label_max_repeat"] == 3  # ตัวนับยังนับครบ ไม่ถูกยุบ


def test_repeated_header_in_middle_is_skipped_and_counted(tmp_path):
    wb = Workbook()
    ws = wb.active
    simple_sheet(ws, rows=[(1, "ก", 1), (2, "ข", 2)])
    for j, h in enumerate(HEAD, 1):
        ws.cell(5, j, h)
    for j, v in enumerate([3, "ค", 3], 1):
        ws.cell(6, j, v)
    (u,) = extract(save(wb, tmp_path)).units
    assert "รหัส: รหัส" not in u.text
    assert u.signals["xlsx:repeated_header_rows"] == 1
    assert u.signals["xlsx:data_rows"] == 3


def test_repeated_two_level_header_skips_both_rows(tmp_path):
    wb = Workbook()
    ws = wb.active
    for top in (1, 5):
        for col, name in zip("ABC", HEAD):
            ws[f"{col}{top}"] = name
            ws.merge_cells(f"{col}{top}:{col}{top + 1}")
    ws["A3"], ws["B3"], ws["C3"] = 1, "ก", 1
    ws["A7"], ws["B7"], ws["C7"] = 2, "ข", 2
    (u,) = extract(save(wb, tmp_path)).units
    assert u.signals["xlsx:repeated_header_rows"] == 1 and u.signals["xlsx:data_rows"] == 2


# ---------- ป้ายกลุ่ม / หมายเหตุ / ค่าในแถว ----------

def test_group_label_and_note_lines(tmp_path):
    wb = Workbook()
    ws = wb.active
    simple_sheet(ws, rows=[])
    ws.cell(2, 2, "ฝ่าย ก")
    ws.cell(3, 1, 1), ws.cell(3, 2, "คน1"), ws.cell(3, 3, 1)
    ws.cell(4, 3, "ข้อมูล ณ 14 พ.ค.")  # ค่าเดียวในคอลัมน์ 3
    (u,) = extract(save(wb, tmp_path)).units
    assert lines(u) == ["## ฝ่าย ก", "รหัส: 1 | ชื่อ: คน1 | เบอร์ต่อ: 1", "หมายเหตุ: ข้อมูล ณ 14 พ.ค."]
    assert [b.type for b in u.blocks] == [BlockType.PARAGRAPH, BlockType.TABLE, BlockType.PARAGRAPH]


def test_label_in_column_1_also_counts(tmp_path):
    wb = Workbook()
    ws = wb.active
    simple_sheet(ws, rows=[(1, "ก", 1)])
    ws["A3"] = "หมวด ข"
    (u,) = extract(save(wb, tmp_path)).units
    assert lines(u)[-1] == "## หมวด ข"


def test_dash_is_kept_blank_and_whitespace_dropped(tmp_path):
    wb = Workbook()
    ws = wb.active
    simple_sheet(
        ws, header=["ลำดับ", "ตำแหน่ง", "ทั้งหมด", "ปัจจุบัน"], rows=[(1, "นิติกร", 1, " -"), (2, "ครู", 2, "  ")]
    )
    (u,) = extract(save(wb, tmp_path)).units
    assert lines(u) == [
        "ลำดับ: 1 | ตำแหน่ง: นิติกร | ทั้งหมด: 1 | ปัจจุบัน: -",
        "ลำดับ: 2 | ตำแหน่ง: ครู | ทั้งหมด: 2",
    ]


def test_merged_cell_in_data_fills_every_covered_row(tmp_path):
    wb = Workbook()
    ws = wb.active
    simple_sheet(ws, header=["ฝ่าย", "ชื่อ", "ต่อ"], rows=[("บัญชี", "ก", 1), (None, "ข", 2), (None, "ค", 3)])
    ws.merge_cells("A2:A4")
    (u,) = extract(save(wb, tmp_path)).units
    assert lines(u) == [
        "ฝ่าย: บัญชี | ชื่อ: ก | ต่อ: 1",
        "ฝ่าย: บัญชี | ชื่อ: ข | ต่อ: 2",
        "ฝ่าย: บัญชี | ชื่อ: ค | ต่อ: 3",
    ]
    assert u.signals["xlsx:merged_cells"] == 1
    (m,) = u.blocks[0].table.merged_cells
    assert (m.row, m.col, m.rowspan, m.colspan) == (1, 0, 3, 1)


def test_headerless_table_text_only(tmp_path):
    wb = Workbook()
    ws = wb.active
    ws["A1"] = "เฉพาะข้อความ"
    ws["A2"] = "บรรทัดสอง"
    (u,) = extract(save(wb, tmp_path)).units
    assert lines(u) == ["เฉพาะข้อความ", "บรรทัดสอง"]
    assert all(b.type == BlockType.PARAGRAPH for b in u.blocks)


# ---------- แบ่ง unit + ป้ายก่อนหน้า ----------

def test_split_unit_adds_previous_labels_context(tmp_path, monkeypatch):
    monkeypatch.setattr(xlsx_mod, "MAX_UNIT_CHARS", 260)
    wb = Workbook()
    ws = wb.active
    ws["A1"] = "ชื่อเรื่องตาราง"
    for j, h in enumerate(HEAD, 1):
        ws.cell(2, j, h)
    r, n = 3, 0
    for g in ["หน่วย ก", "ประเภท ข", "ประเภท ค", "หน่วย ง", "ประเภท จ"]:
        ws.cell(r, 2, g)
        r += 1
        for _ in range(3):
            n += 1
            ws.cell(r, 1, n), ws.cell(r, 2, f"พนักงานคนที่{n}"), ws.cell(r, 3, n)
            r += 1
    d = extract(save(wb, tmp_path))
    assert len(d.units) > 1
    first, rest = d.units[0], d.units[1:]
    assert not any(ln.startswith("[ป้ายก่อนหน้า") for ln in lines(first))
    for u in rest:
        assert lines(u)[0] == "ชื่อเรื่องตาราง"
        assert lines(u)[1].startswith("[ป้ายก่อนหน้า: ")
        assert u.blocks[1].type == BlockType.PARAGRAPH
        assert "ชื่อเรื่องตาราง" in u.label
    for u in d.units:  # ป้ายท้าย unit ไม่ถูกทิ้งไว้ลำพัง
        assert not lines(u)[-1].startswith("## ")
    got = [ln for u in d.units for ln in lines(u) if ln.startswith("รหัส:")]
    assert got == [f"รหัส: {k} | ชื่อ: พนักงานคนที่{k} | เบอร์ต่อ: {k}" for k in range(1, 16)]
    assert sum(u.signals["xlsx:data_rows"] for u in d.units) == 15


def test_context_is_last_three_labels_in_order_of_appearance(tmp_path, monkeypatch):
    monkeypatch.setattr(xlsx_mod, "MAX_UNIT_CHARS", 200)
    wb = Workbook()
    ws = wb.active
    for j, h in enumerate(HEAD, 1):
        ws.cell(1, j, h)
    r = 2
    for g in ["ก", "ข", "ก", "ค"]:
        ws.cell(r, 2, g)
        r += 1
    for k in range(1, 8):
        ws.cell(r, 1, k), ws.cell(r, 2, "คนที่ " + "x" * 30), ws.cell(r, 3, k)
        r += 1
    d = extract(save(wb, tmp_path))
    ctx = [ln for ln in lines(d.units[1]) if ln.startswith("[ป้ายก่อนหน้า")]
    assert ctx == ["[ป้ายก่อนหน้า: ข → ก → ค]"]
    assert d.units[0].signals["xlsx:group_label_max_repeat"] == 2


# ---------- ค่าในเซลล์ ----------

def test_value_formatting(tmp_path):
    wb = Workbook()
    ws = wb.active
    ws.append(["a", "b", "c", "d", "e", "f", "g"])
    ws.append([3.0, 0.185, dt.datetime(2025, 5, 14), dt.datetime(2025, 5, 14, 9, 30), 2.5, True, "#N/A"])
    ws["B2"].number_format = "0.0%"
    (u,) = extract(save(wb, tmp_path)).units
    assert lines(u) == ["a: 3 | b: 18.5% | c: 2025-05-14 | d: 2025-05-14 09:30 | e: 2.5 | f: TRUE | g: #N/A"]


def test_leading_zero_text_is_preserved(tmp_path):
    wb = Workbook()
    ws = wb.active
    ws.append(["รหัสไปรษณีย์", "เบอร์"])
    ws.append(["01000", "0812345678"])
    (u,) = extract(save(wb, tmp_path)).units
    assert lines(u) == ["รหัสไปรษณีย์: 01000 | เบอร์: 0812345678"]


def test_formula_without_cached_value_is_blank_and_counted(tmp_path):
    wb = Workbook()
    ws = wb.active
    ws.append(["ก", "ข", "รวม"])
    ws.append([1, 2, "=A2+B2"])  # openpyxl ไม่เก็บค่าที่คำนวณ
    (u,) = extract(save(wb, tmp_path)).units
    assert lines(u) == ["ก: 1 | ข: 2"]
    assert u.signals["xlsx:formula_no_cache"] == 1


def test_normalizer_applies_to_cells(tmp_path):
    wb = Workbook()
    ws = wb.active
    ws.append(["ชื่อ", "งาน"])
    ws.append(["สมชาย", "ธุุรกิิจ ดําเนินงาน"])
    (u,) = extract(save(wb, tmp_path)).units
    assert "งาน: ธุรกิจ ดำเนินงาน" in u.text
    assert u.quality.normalization_changes > 0


# ---------- ไฟล์ใหญ่ ----------

def test_sheet_xml_size_gate(tmp_path, monkeypatch):
    wb = Workbook()
    simple_sheet(wb.active)
    f = save(wb, tmp_path)
    monkeypatch.setattr(xlsx_mod, "MAX_SHEET_XML_BYTES", 100)
    d = extract_file(f)
    assert d.status == DocumentStatus.FAILED and d.error.code == ErrorCode.TOO_LARGE
    assert d.error.message.startswith("ไฟล์เป็นฐานข้อมูลขนาดใหญ่ ไม่เหมาะกับการค้นแบบเอกสาร")
    assert d.detected_type == DetectedType.XLSX


def test_sheet_xml_gate_does_not_open_workbook(tmp_path, monkeypatch):
    wb = Workbook()
    simple_sheet(wb.active)
    f = save(wb, tmp_path)
    monkeypatch.setattr(xlsx_mod, "MAX_SHEET_XML_BYTES", 100)

    def boom(*a, **k):
        raise AssertionError("ต้องปฏิเสธก่อนเปิดด้วย openpyxl")

    monkeypatch.setattr(xlsx_mod.openpyxl, "load_workbook", boom)
    assert extract_file(f).error.code == ErrorCode.TOO_LARGE


def test_row_count_gate_with_write_only_file(tmp_path):
    wb = Workbook(write_only=True)
    ws = wb.create_sheet()
    ws.append(["a", "b"])
    for i in range(xlsx_mod.MAX_DATA_ROWS):  # หัว + MAX_DATA_ROWS แถว = เกินเกณฑ์หนึ่งแถว
        ws.append([i, "x"])
    f = tmp_path / "big.xlsx"
    wb.save(f)
    d = extract_file(f)
    assert d.error is not None and d.error.code == ErrorCode.TOO_LARGE
    assert "50,000" in d.error.message


def test_row_count_at_limit_is_accepted(tmp_path, monkeypatch):
    monkeypatch.setattr(xlsx_mod, "MAX_DATA_ROWS", 5)
    wb = Workbook()
    ws = wb.active
    ws.append(["a", "b"])
    for i in range(4):
        ws.append([i, "x"])
    assert extract_file(save(wb, tmp_path)).error is None  # 5 แถวพอดี
    ws.append([9, "x"])
    assert extract_file(save(wb, tmp_path, "b.xlsx")).error.code == ErrorCode.TOO_LARGE


def test_blank_styled_rows_do_not_count_toward_row_limit(tmp_path, monkeypatch):
    monkeypatch.setattr(xlsx_mod, "MAX_DATA_ROWS", 5)
    wb = Workbook()
    ws = wb.active
    simple_sheet(ws)
    for r in range(10, 40):
        ws.cell(r, 1).fill = PatternFill("solid", fgColor="FFFF00")
    assert extract_file(save(wb, tmp_path)).error is None


def test_corrupt_workbook_is_corrupt_error(tmp_path):
    f = tmp_path / "a.xlsx"
    with zipfile.ZipFile(f, "w") as z:
        z.writestr("xl/workbook.xml", "not xml")
    d = extract_file(f)
    assert d.detected_type == DetectedType.XLSX and d.error.code == ErrorCode.CORRUPT
