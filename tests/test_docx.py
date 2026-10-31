"""DOCX extractor (app/extraction/docx.py) ด้วยไฟล์สังเคราะห์ (python-docx หรือ XML เขียนมือ)"""

import io
import zipfile

import pymupdf
import pytest
from docx import Document as WordDocument
from docx.shared import Pt

from app.extraction.docx import MAX_UNIT_CHARS, THAI_LETTERS, format_number
from app.extraction.pipeline import extract_file
from app.models import BlockType, DetectedType, DocumentStatus, ErrorCode, UnitStatus, UnitType

BODY = Pt(12)


# ---------- ตัวช่วยสร้างไฟล์ ----------

def para(doc, text="", bold=False, size=BODY, style=None):
    p = doc.add_paragraph(style=style)
    if text:
        r = p.add_run(text)
        r.bold = bold
        r.font.size = size
    return p


def body(doc, n=3, text="เนื้อหาปกติของเอกสารสำหรับทดสอบการอ่านไฟล์"):
    for _ in range(n):
        para(doc, text)


def save(doc, tmp_path, name="a.docx"):
    f = tmp_path / name
    doc.save(str(f))
    return f


def extract(path):
    d = extract_file(path)
    assert d.error is None, d.error
    return d


def paths(d):
    return [u.position["heading_path"] for u in d.units]


NS_DECL = (
    'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" '
    'xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006" '
    'xmlns:wps="http://schemas.microsoft.com/office/word/2010/wordprocessingShape" '
    'xmlns:wp="http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing" '
    'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" '
    'xmlns:v="urn:schemas-microsoft-com:vml" '
    'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships"'
)


def raw_docx(path, body_xml, numbering=None, footnotes=None, styles=None):
    """DOCX ขั้นต่ำที่เขียน XML เอง (ไม่ผ่าน python-docx) ไว้ทดสอบ numbering / กล่องข้อความ / เชิงอรรถ"""
    with zipfile.ZipFile(path, "w") as z:
        z.writestr("[Content_Types].xml", '<?xml version="1.0"?><Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types"/>')
        z.writestr("word/document.xml", f'<?xml version="1.0" encoding="UTF-8"?><w:document {NS_DECL}><w:body>{body_xml}</w:body></w:document>')
        if numbering:
            z.writestr("word/numbering.xml", f'<?xml version="1.0" encoding="UTF-8"?><w:numbering {NS_DECL}>{numbering}</w:numbering>')
        if footnotes:
            z.writestr("word/footnotes.xml", f'<?xml version="1.0" encoding="UTF-8"?><w:footnotes {NS_DECL}>{footnotes}</w:footnotes>')
        if styles:
            z.writestr("word/styles.xml", f'<?xml version="1.0" encoding="UTF-8"?><w:styles {NS_DECL}>{styles}</w:styles>')
    return path


def p_xml(text, num=None, style=None, bold=False, size=24):
    ppr = ""
    if style:
        ppr += f'<w:pStyle w:val="{style}"/>'
    if num is not None:
        ppr += f'<w:numPr><w:ilvl w:val="{num[1]}"/><w:numId w:val="{num[0]}"/></w:numPr>'
    rpr = f'<w:rPr>{"<w:b/>" if bold else ""}<w:sz w:val="{size}"/></w:rPr>'
    return f"<w:p><w:pPr>{ppr}</w:pPr><w:r>{rpr}<w:t>{text}</w:t></w:r></w:p>"


def lvl(i, start, fmt, text):
    return f'<w:lvl w:ilvl="{i}"><w:start w:val="{start}"/><w:numFmt w:val="{fmt}"/><w:lvlText w:val="{text}"/></w:lvl>'


def abstract(aid, *levels):
    return f'<w:abstractNum w:abstractNumId="{aid}">{"".join(levels)}</w:abstractNum>'


def num(nid, aid, extra=""):
    return f'<w:num w:numId="{nid}"><w:abstractNumId w:val="{aid}"/>{extra}</w:num>'


def lines(d):
    return [ln for u in d.units for ln in u.text.split("\n")]


# ---------- หัวข้อ ----------

def test_heading_rules(tmp_path):
    doc = WordDocument()
    para(doc, "ภาพรวมของระบบ", bold=True, size=Pt(20))  # ใหญ่กว่า body + 4pt -> ระดับ 1
    body(doc)
    para(doc, "รายละเอียดการทำงาน", bold=True, size=Pt(14))  # ใหญ่กว่า body แต่ไม่ถึง +4 -> ระดับ 2
    body(doc)
    para(doc, "1.2.3 หัวข้อย่อย", bold=True, size=BODY)  # ขนาดเท่าเนื้อหา แต่มีเลข N.N.N -> ระดับ 3
    body(doc)
    para(doc, "ข้อสำคัญ", bold=True, size=BODY)  # ตัวหนาสั้นขนาดเท่าเนื้อหา ไม่มีเลข -> ไม่ใช่หัวข้อ
    para(doc, "1.2 รายการที่ไม่หนา", bold=False, size=Pt(14))  # ใหญ่แต่ไม่หนา -> ไม่ใช่หัวข้อ
    para(doc, "ข้อความตัวหนายาวเกินแปดสิบตัวอักษร " * 4, bold=True, size=Pt(20))  # ยาวเกิน 80
    d = extract(save(doc, tmp_path))
    assert paths(d) == [
        ["ภาพรวมของระบบ", "รายละเอียดการทำงาน", "1.2.3 หัวข้อย่อย"],
    ] or paths(d)[-1] == ["ภาพรวมของระบบ", "รายละเอียดการทำงาน", "1.2.3 หัวข้อย่อย"]
    last = d.units[-1]
    texts = [b.text for b in last.blocks]
    assert "ข้อสำคัญ" in texts and "1.2 รายการที่ไม่หนา" in texts
    assert all(b.type != BlockType.HEADING for b in last.blocks[1:])
    levels = {u.position["heading_path"][-1]: u.blocks[0].level for u in d.units}
    assert levels == {"ภาพรวมของระบบ": 1, "รายละเอียดการทำงาน": 2, "1.2.3 หัวข้อย่อย": 3}
    assert d.units[0].signals["docx:heading_by_format"] == 1.0
    assert d.units[2].signals["docx:heading_by_numbering"] == 1.0


def test_heading_style_is_only_a_signal(tmp_path):
    doc = WordDocument()
    para(doc, "ตัวอย่างหัวข้อ style", bold=False, size=BODY, style="Heading 1")  # style Heading แต่ไม่หนา/ไม่ใหญ่
    body(doc)
    d = extract(save(doc, tmp_path))
    assert paths(d) == [[]]  # ไม่ถูกตัดสินเป็นหัวข้อ
    doc = WordDocument()
    para(doc, "หัวข้อที่หนาและใหญ่", bold=True, size=Pt(20), style="Heading 1")
    body(doc)
    u = extract(save(doc, tmp_path, "b.docx")).units[0]
    assert u.signals["docx:heading_style"] == 1.0 and u.signals["docx:heading_by_format"] == 1.0


def test_chapter_and_title_are_merged(tmp_path):
    doc = WordDocument()
    para(doc, "บทที่ 1", bold=True, size=Pt(20))
    para(doc, "บทนำ", bold=True, size=Pt(24))
    body(doc)
    para(doc, "1.1 ความเป็นมา", bold=True, size=Pt(14))
    body(doc)
    para(doc, "บทที่ 2", bold=True, size=Pt(20))
    para(doc, "2.1 หัวข้อที่มีเลข", bold=True, size=Pt(14))  # ย่อหน้าถัดไปมีเลข -> ไม่รวม
    body(doc)
    para(doc, "บทที่ 3 วิธีการ", bold=True, size=Pt(20))  # อยู่ในย่อหน้าเดียวแล้ว
    body(doc)
    d = extract(save(doc, tmp_path))
    assert paths(d) == [
        ["บทที่ 1 บทนำ"],
        ["บทที่ 1 บทนำ", "1.1 ความเป็นมา"],
        ["บทที่ 2", "2.1 หัวข้อที่มีเลข"],
        ["บทที่ 3 วิธีการ"],
    ]
    assert d.units[0].blocks[0].level == 1 and d.units[0].blocks[0].text == "บทที่ 1 บทนำ"


def test_appendix_shifts_levels(tmp_path):
    doc = WordDocument()
    para(doc, "บทที่ 1 บทนำ", bold=True, size=Pt(20))
    body(doc)
    para(doc, "ภาคผนวก ก.", bold=True, size=Pt(20))
    para(doc, "บทที่ 1 บทนำ", bold=True, size=Pt(20))  # ซ้ำกับของจริง
    body(doc)
    para(doc, "1.1 หัวข้อในภาคผนวก", bold=True, size=Pt(14))
    body(doc)
    para(doc, "ภาคผนวก ข.", bold=True, size=Pt(20))
    body(doc)
    d = extract(save(doc, tmp_path))
    assert paths(d) == [
        ["บทที่ 1 บทนำ"],
        ["ภาคผนวก ก.", "บทที่ 1 บทนำ"],
        ["ภาคผนวก ก.", "บทที่ 1 บทนำ", "1.1 หัวข้อในภาคผนวก"],
        ["ภาคผนวก ข."],
    ]
    assert [u.blocks[0].level for u in d.units] == [1, 2, 3, 1]


def test_heading_without_content_makes_no_unit_but_stays_in_path(tmp_path):
    doc = WordDocument()
    para(doc, "บทที่ 1 บทนำ", bold=True, size=Pt(20))
    para(doc, "1.1 หัวข้อลูก", bold=True, size=Pt(14))
    body(doc)
    para(doc, "1.2 หัวข้อว่าง", bold=True, size=Pt(14))
    para(doc, "1.3 หัวข้อสุดท้าย", bold=True, size=Pt(14))
    d = extract(save(doc, tmp_path))
    assert paths(d) == [["บทที่ 1 บทนำ", "1.1 หัวข้อลูก"]]


def test_preface_before_first_heading(tmp_path):
    doc = WordDocument()
    para(doc, "ข้อความก่อนหัวข้อแรกของเอกสาร")
    para(doc, "บทที่ 1 บทนำ", bold=True, size=Pt(20))
    body(doc)
    d = extract(save(doc, tmp_path))
    assert d.units[0].label == "ส่วนต้นเอกสาร" and d.units[0].position["heading_path"] == []
    assert d.units[0].text == "ข้อความก่อนหัวข้อแรกของเอกสาร"
    assert d.units[1].label == "บทที่ 1 บทนำ"


# ---------- การแบ่ง unit / label ----------

def test_split_at_limit_keeps_heading_and_paragraph_boundaries(tmp_path):
    doc = WordDocument()
    para(doc, "บทที่ 1 บทนำ", bold=True, size=Pt(20))
    for i in range(9):
        para(doc, f"ย่อหน้า{i} " + "ก" * 1290)
    d = extract(save(doc, tmp_path))
    assert len(d.units) == 3
    for k, u in enumerate(d.units, 1):
        assert len(u.text) <= MAX_UNIT_CHARS
        assert u.text.startswith("บทที่ 1 บทนำ\n")
        assert u.label == f"บทที่ 1 บทนำ (ส่วน {k}/3)"
        assert u.position["heading_path"] == ["บทที่ 1 บทนำ"]
        assert u.type == UnitType.SECTION
    ranges = [u.position["paragraphs"] for u in d.units]
    assert ranges[0] == [0, 3] and ranges[1][0] == 4 and ranges[2][1] == 9
    assert sum(u.text.count("ย่อหน้า") for u in d.units) == 9  # ไม่ซ้ำ ไม่หาย


def test_single_long_paragraph_is_never_cut(tmp_path):
    doc = WordDocument()
    para(doc, "บทที่ 1 บทนำ", bold=True, size=Pt(20))
    para(doc, "สั้น")
    para(doc, "ข" * 5000)
    para(doc, "ท้าย")
    d = extract(save(doc, tmp_path))
    big = [u for u in d.units if "ข" * 5000 in u.text]
    assert len(big) == 1 and len(big[0].text) > MAX_UNIT_CHARS
    assert [u.label for u in d.units] == ["บทที่ 1 บทนำ (ส่วน 1/3)", "บทที่ 1 บทนำ (ส่วน 2/3)", "บทที่ 1 บทนำ (ส่วน 3/3)"]


def test_duplicate_labels_get_suffix(tmp_path):
    doc = WordDocument()
    for _ in range(3):
        para(doc, "1.1 หัวข้อซ้ำ", bold=True, size=Pt(14))
        body(doc, 1)
    d = extract(save(doc, tmp_path))
    assert [u.label for u in d.units] == ["1.1 หัวข้อซ้ำ", "1.1 หัวข้อซ้ำ (ช่วงที่ 2)", "1.1 หัวข้อซ้ำ (ช่วงที่ 3)"]


def test_label_is_full_path(tmp_path):
    doc = WordDocument()
    para(doc, "บทที่ 1 บทนำ", bold=True, size=Pt(20))
    para(doc, "1.2 หัวข้อกลาง", bold=True, size=Pt(14))
    para(doc, "1.2.1 หัวข้อลึก", bold=True, size=BODY)
    body(doc)
    d = extract(save(doc, tmp_path))
    assert d.units[0].label == "บทที่ 1 บทนำ › 1.2 หัวข้อกลาง › 1.2.1 หัวข้อลึก"


# ---------- ตาราง / ย่อหน้า ----------

def test_table_with_merged_cells(tmp_path):
    doc = WordDocument()
    para(doc, "บทที่ 1 บทนำ", bold=True, size=Pt(20))
    para(doc, "ก่อนตาราง")
    t = doc.add_table(rows=3, cols=3)
    for r in range(3):
        for c in range(3):
            t.cell(r, c).text = f"ช่อง{r}{c}"
    t.cell(0, 0).merge(t.cell(0, 1))  # แนวนอน
    t.cell(1, 2).merge(t.cell(2, 2))  # แนวตั้ง
    para(doc, "หลังตาราง")
    d = extract(save(doc, tmp_path))
    blocks = d.units[0].blocks
    tb = [b for b in blocks if b.type == BlockType.TABLE]
    assert len(tb) == 1
    tbl = tb[0].table
    assert len(tbl.rows) == 3 and all(len(r) == 3 for r in tbl.rows)
    assert {(m.row, m.col, m.rowspan, m.colspan) for m in tbl.merged_cells} == {(0, 0, 1, 2), (1, 2, 2, 1)}
    text_lines = tb[0].text.split("\n")
    assert len(text_lines) == 3
    assert text_lines[0].count("ช่อง00") == 1 and text_lines[0].count("|") == 1  # ผสานแล้วไม่ซ้ำค่า
    assert "ช่อง10 | ช่อง11 | ช่อง12 ช่อง22" == text_lines[1]  # python-docx รวมข้อความของเซลล์ที่ผสานแนวตั้งไว้ในเซลล์บน
    assert "ช่อง20 | ช่อง21" == text_lines[2]  # แถวล่างเซลล์ที่ผสานไปแล้วไม่ถูกนับ
    order = [b.text.split("\n")[0] for b in blocks]
    assert order.index("ก่อนตาราง") < order.index(text_lines[0]) < order.index("หลังตาราง")


def test_table_is_one_element_in_paragraph_index(tmp_path):
    doc = WordDocument()
    para(doc, "บทที่ 1 บทนำ", bold=True, size=Pt(20))  # 0
    para(doc, "ก่อน")  # 1
    doc.add_table(rows=1, cols=2).cell(0, 0).text = "ก"  # 2
    para(doc, "หลัง")  # 3
    d = extract(save(doc, tmp_path))
    assert d.units[0].position["paragraphs"] == [0, 3]


def test_line_break_in_paragraph_is_newline_not_replacement_char(tmp_path):
    doc = WordDocument()
    p = doc.add_paragraph()
    r = p.add_run("บรรทัดแรก")
    r.font.size = BODY
    r.add_break()
    p.add_run("บรรทัดสอง").font.size = BODY
    d = extract(save(doc, tmp_path))
    assert d.units[0].text == "บรรทัดแรก\nบรรทัดสอง"
    assert "�" not in d.units[0].text
    assert "norm_sig:fffd_remaining" not in d.units[0].signals


def test_vertical_tab_text_in_run_is_newline():
    from app.thai.normalizer import normalize

    assert normalize("บรรทัดแรก\x0bบรรทัดสอง") == "บรรทัดแรก\nบรรทัดสอง"


def test_images_are_counted(tmp_path):
    png = pymupdf.Pixmap(pymupdf.csRGB, pymupdf.IRect(0, 0, 8, 8)).tobytes("png")
    doc = WordDocument()
    para(doc, "บทที่ 1 บทนำ", bold=True, size=Pt(20))
    para(doc, "มีรูป")
    doc.add_picture(io.BytesIO(png))
    doc.add_picture(io.BytesIO(png))
    d = extract(save(doc, tmp_path))
    assert d.units[0].signals["docx:image_count"] == 2.0


def test_header_and_footer_are_not_read(tmp_path):
    doc = WordDocument()
    doc.sections[0].header.paragraphs[0].text = "ชื่อเอกสารซ้ำทุกหน้า"
    doc.sections[0].footer.paragraphs[0].text = "เลขหน้า"
    para(doc, "เนื้อหา")
    d = extract(save(doc, tmp_path))
    assert "ชื่อเอกสารซ้ำทุกหน้า" not in d.units[0].text and "เลขหน้า" not in d.units[0].text


# ---------- เลขอัตโนมัติ ----------

def run_numbering(tmp_path, numbering, paragraphs, styles=None):
    f = raw_docx(tmp_path / "n.docx", "".join(paragraphs), numbering=numbering, styles=styles)
    return lines(extract(f))


def test_multilevel_numbering_like_word(tmp_path):
    # แบบเดียวกับไฟล์จริง: ระดับ 0 เริ่ม 3, ระดับ 1 เริ่ม 1, lvlText "%1.%2"
    numbering = abstract(0, lvl(0, 3, "decimal", "%1"), lvl(1, 1, "decimal", "%1.%2"), lvl(2, 1, "decimal", "%1.%2.%3")) + num(1, 0)
    ps = [p_xml("A", (1, 1)), p_xml("B", (1, 1)), p_xml("C", (1, 0)), p_xml("D", (1, 1)), p_xml("E", (1, 2)),
          p_xml("F", (1, 0)), p_xml("G", (1, 1))]
    assert run_numbering(tmp_path, numbering, ps) == ["3.1 A", "3.2 B", "3 C", "3.1 D", "3.1.1 E", "4 F", "4.1 G"]


def test_numbering_start_values_on_unused_parent_level(tmp_path):
    # เหมือน "ส่วนที่ 2 ส่วนกลาง" -> 2.2 : ระดับบนไม่เคยใช้ แสดงค่า start ของมัน
    numbering = abstract(0, lvl(0, 2, "decimal", "%1"), lvl(1, 2, "decimal", "%1.%2")) + num(1, 0)
    assert run_numbering(tmp_path, numbering, [p_xml("ส่วนที่ 2 ส่วนกลาง", (1, 1))]) == ["2.2 ส่วนที่ 2 ส่วนกลาง"]


def test_numbering_formats(tmp_path):
    numbering = (
        abstract(0, lvl(0, 1, "thaiLetters", "%1."), lvl(1, 1, "lowerRoman", "(%2)"))
        + abstract(1, lvl(0, 1, "upperLetter", "%1)"), lvl(1, 1, "lowerLetter", "%2."))
        + abstract(2, lvl(0, 1, "upperRoman", "%1"), lvl(1, 1, "thaiNumbers", "%2 "))
        + num(1, 0) + num(2, 1) + num(3, 2)
    )
    ps = [p_xml("ก", (1, 0)), p_xml("ข", (1, 0)), p_xml("ย่อย", (1, 1)), p_xml("ย่อย", (1, 1)),
          p_xml("x", (2, 0)), p_xml("y", (2, 0)), p_xml("z", (2, 1)),
          p_xml("r", (3, 0)), p_xml("r", (3, 0)), p_xml("r", (3, 0)), p_xml("t", (3, 1)), p_xml("t", (3, 1))]
    assert run_numbering(tmp_path, numbering, ps) == [
        "ก. ก", "ข. ข", "(i) ย่อย", "(ii) ย่อย",
        "A) x", "B) y", "a. z",
        "I r", "II r", "III r", "1 t", "2 t",  # thaiNumbers "๑" ถูก normalize เป็น 1
    ]


def test_format_number_helpers():
    assert format_number(1, "thaiLetters") == "ก" and format_number(3, "thaiLetters") == "ค"
    assert format_number(len(THAI_LETTERS) + 1, "thaiLetters") == "กก"
    assert format_number(27, "lowerLetter") == "aa" and format_number(4, "upperRoman") == "IV"
    assert format_number(1994, "lowerRoman") == "mcmxciv" and format_number(7, "decimalZero") == "07"
    assert format_number(12, "unknownFormat") == "12"


def test_bullet_has_no_number_but_is_a_list(tmp_path):
    numbering = abstract(0, lvl(0, 1, "bullet", "•")) + num(1, 0)
    f = raw_docx(tmp_path / "b.docx", p_xml("รายการหัวข้อย่อย", (1, 0)) + p_xml("ปกติ"), numbering=numbering)
    d = extract(f)
    assert [b.type for b in d.units[0].blocks] == [BlockType.LIST, BlockType.PARAGRAPH]
    assert d.units[0].text == "รายการหัวข้อย่อย\nปกติ"


def test_start_override_and_num_zero(tmp_path):
    numbering = (
        abstract(0, lvl(0, 1, "decimal", "%1."))
        + num(1, 0)
        + num(2, 0, '<w:lvlOverride w:ilvl="0"><w:startOverride w:val="10"/></w:lvlOverride>')
    )
    ps = [p_xml("a", (1, 0)), p_xml("b", (1, 0)), p_xml("c", (2, 0)), p_xml("d", (2, 0)), p_xml("ไม่มีเลข", (0, 0))]
    assert run_numbering(tmp_path, numbering, ps) == ["1. a", "2. b", "10. c", "11. d", "ไม่มีเลข"]


def test_numbering_from_paragraph_style(tmp_path):
    numbering = abstract(0, lvl(0, 1, "decimal", "ข้อ %1")) + num(1, 0)
    styles = '<w:style w:type="paragraph" w:styleId="NumStyle"><w:name w:val="NumStyle"/><w:pPr><w:numPr><w:numId w:val="1"/></w:numPr></w:pPr></w:style>'
    ps = [p_xml("หนึ่ง", style="NumStyle"), p_xml("สอง", style="NumStyle")]
    assert run_numbering(tmp_path, numbering, ps, styles=styles) == ["ข้อ 1 หนึ่ง", "ข้อ 2 สอง"]


def test_auto_number_makes_heading(tmp_path):
    # ข้อความไม่มีเลข แต่เลขอัตโนมัติ "2.2" + ตัวหนาใหญ่กว่า body => หัวข้อระดับ 2 (ตัดสินจาก numbering)
    numbering = abstract(0, lvl(0, 2, "decimal", "%1"), lvl(1, 2, "decimal", "%1.%2")) + num(1, 0)
    ps = [p_xml("เนื้อหา", size=24)] * 3 + [p_xml("ส่วนกลาง", (1, 1), bold=True, size=28), p_xml("ข้อความใต้หัวข้อ", size=24)]
    d = extract(raw_docx(tmp_path / "h.docx", "".join(ps), numbering=numbering))
    assert paths(d)[-1] == ["2.2 ส่วนกลาง"]
    assert d.units[-1].signals["docx:heading_by_numbering"] == 1.0


# ---------- กล่องข้อความ / เชิงอรรถ ----------

def test_textbox_follows_host_paragraph_once(tmp_path):
    box = '<w:p><w:r><w:t>ข้อความในกล่อง</w:t></w:r></w:p>'
    host = (
        "<w:p><w:r><w:t>ย่อหน้าที่มีกล่อง</w:t></w:r><w:r>"
        '<mc:AlternateContent><mc:Choice Requires="wps"><w:drawing><wp:inline><a:graphic><a:graphicData>'
        f"<wps:wsp><wps:txbx><w:txbxContent>{box}</w:txbxContent></wps:txbx></wps:wsp>"
        "</a:graphicData></a:graphic></wp:inline></w:drawing></mc:Choice>"
        f"<mc:Fallback><w:pict><v:shape><v:textbox><w:txbxContent>{box}</w:txbxContent></v:textbox></v:shape></w:pict></mc:Fallback>"
        "</mc:AlternateContent></w:r></w:p>"
    )
    d = extract(raw_docx(tmp_path / "t.docx", p_xml("ก่อน") + host + p_xml("หลัง")))
    assert d.units[0].text == "ก่อน\nย่อหน้าที่มีกล่อง\nข้อความในกล่อง\nหลัง"


def test_footnote_is_appended_to_unit(tmp_path):
    footnotes = (
        '<w:footnote w:type="separator" w:id="0"><w:p><w:r><w:separator/></w:r></w:p></w:footnote>'
        '<w:footnote w:id="2"><w:p><w:r><w:t>ที่มาของข้อมูล</w:t></w:r></w:p></w:footnote>'
    )
    host = '<w:p><w:r><w:t>ข้อความอ้างอิง</w:t></w:r><w:r><w:footnoteReference w:id="2"/></w:r></w:p>'
    d = extract(raw_docx(tmp_path / "f.docx", host + p_xml("ย่อหน้าถัดไป"), footnotes=footnotes))
    assert d.units[0].text == "ข้อความอ้างอิง\nย่อหน้าถัดไป\nเชิงอรรถ: ที่มาของข้อมูล"
    assert d.units[0].blocks[-1].type == BlockType.PARAGRAPH


# ---------- pipeline / ด่านตรวจ / error ----------

def test_docx_goes_through_pipeline_and_gate(tmp_path):
    doc = WordDocument()
    para(doc, "บทที่ 1 บทนำ", bold=True, size=Pt(20))
    para(doc, "ธนาคารมีนโยบายกำกับดูแลกิจการที่ดี " * 6)
    para(doc, "1.1 หัวข้อสั้น", bold=True, size=Pt(14))
    para(doc, "สั้น")
    d = extract(save(doc, tmp_path))
    assert d.detected_type == DetectedType.DOCX and d.status == DocumentStatus.DONE
    assert d.pipeline_version and d.has_outline is True
    first, second = d.units
    assert first.status == UnitStatus.OK and not first.quality.flags
    assert second.status == UnitStatus.OK and [f.value for f in second.quality.flags] == ["low_text"]  # ไม่มี image_coverage -> ไม่มี scan_like
    assert "gate:thai_chars" in first.signals and "image_coverage" not in first.signals
    assert first.signals["char_count"] == len(first.text.strip())
    assert first.method.value == "office_xml"


def test_normalizer_applies_to_docx_text(tmp_path):
    doc = WordDocument()
    para(doc, "ธุุรกิิจ ๒๕๖๘")
    d = extract(save(doc, tmp_path))
    assert d.units[0].text == "ธุรกิจ 2568"
    assert d.units[0].signals["norm_fix:dup_mark"] >= 1 and d.units[0].quality.normalization_changes >= 3


def test_empty_docx_gives_one_empty_unit(tmp_path):
    d = extract(save(WordDocument(), tmp_path))
    assert d.unit_count == 1 and d.units[0].text == "" and d.status == DocumentStatus.DONE
    assert [f.value for f in d.units[0].quality.flags] == ["empty"]


def test_corrupt_docx_reports_error(tmp_path):
    f = tmp_path / "c.docx"
    with zipfile.ZipFile(f, "w") as z:
        z.writestr("word/document.xml", "this is not xml")
    d = extract_file(f)
    assert d.detected_type == DetectedType.DOCX and d.error.code == ErrorCode.CORRUPT and d.status == DocumentStatus.FAILED
    g = tmp_path / "d.docx"
    with zipfile.ZipFile(g, "w") as z:
        z.writestr("word/other.xml", "<a/>")
    assert extract_file(g).error.code == ErrorCode.CORRUPT


def test_childless_sibling_heading_text_is_not_lost(tmp_path):
    doc = WordDocument()
    para(doc, "1.1 หัวข้อพี่น้องที่ไม่มีเนื้อหา", bold=True, size=Pt(14))
    para(doc, "1.2 หัวข้อที่มีเนื้อหา", bold=True, size=Pt(14))
    body(doc, 1)
    para(doc, "1.3 หัวข้อเปล่าตอนท้าย", bold=True, size=Pt(14))
    d = extract(save(doc, tmp_path))
    assert d.unit_count == 1
    u = d.units[0]
    assert u.position["heading_path"] == ["1.2 หัวข้อที่มีเนื้อหา"] and u.text.startswith("1.2 หัวข้อที่มีเนื้อหา\n")
    assert "1.1 หัวข้อพี่น้องที่ไม่มีเนื้อหา" in u.text and "1.3 หัวข้อเปล่าตอนท้าย" in u.text
    assert u.position["paragraphs"] == [0, 3]
    # หัวข้อพ่อที่ไม่มีเนื้อหาไม่ถูกยกซ้ำ (อยู่ใน heading_path แล้ว)
    doc = WordDocument()
    para(doc, "บทที่ 1 บทนำ", bold=True, size=Pt(20))
    para(doc, "1.1 หัวข้อลูก", bold=True, size=Pt(14))
    body(doc, 1)
    d = extract(save(doc, tmp_path, "b.docx"))
    assert "บทที่ 1 บทนำ" not in d.units[0].text
