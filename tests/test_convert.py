"""แปลงไฟล์รุ่นเก่า (app/extraction/convert.py) โดยฉีดคำสั่งแปลงแทน LibreOffice จริง"""

import subprocess
import sys
import textwrap
import time
from pathlib import Path

import pytest
from docx import Document as WordDocument
from openpyxl.workbook import Workbook
from pptx import Presentation

from app.extraction import convert as convert_mod
from app.extraction import pipeline as pipeline_mod
from app.extraction.pipeline import extract_file
from app.models import DetectedType, DocumentStatus, ErrorCode, ExtractionMethod


def make_docx(path: Path):
    d = WordDocument()
    d.add_paragraph("ทดสอบการแปลงไฟล์")
    d.save(str(path))


def make_xlsx(path: Path):
    wb = Workbook()
    wb.active.append(["ก", "ข"])
    wb.active.append([1, 2])
    wb.save(path)


def make_pptx(path: Path):
    prs = Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[5])
    s.shapes.title.text = "สไลด์ทดสอบ"
    prs.save(str(path))


MAKERS = {DetectedType.DOC: ("docx", make_docx), DetectedType.XLS: ("xlsx", make_xlsx), DetectedType.PPT: ("pptx", make_pptx)}


class FakeRun:
    """แทน ``_run``: เก็บคำสั่งไว้ดู แล้วเขียนผลลัพธ์ (หรือไม่เขียน / โยน timeout) ตามที่สั่ง"""

    def __init__(self, kind: DetectedType, mode: str = "ok"):
        self.kind, self.mode = kind, mode
        self.cmd: list[str] = []
        self.timeout = None

    def __call__(self, cmd, timeout):
        self.cmd, self.timeout = cmd, timeout
        if self.mode == "timeout":
            raise subprocess.TimeoutExpired(cmd, timeout)
        out_dir = Path(cmd[cmd.index("--outdir") + 1])
        (out_dir / "input.lock").write_text("ขยะที่ LibreOffice ทิ้งไว้")  # ต้องถูกลบพร้อมโฟลเดอร์
        if self.mode == "ok":
            ext, maker = MAKERS[self.kind]
            maker(out_dir / f"input.{ext}")
        return 1 if self.mode == "fail" else 0

    @property
    def work_dir(self) -> Path:
        return Path(self.cmd[self.cmd.index("--outdir") + 1]).parent


@pytest.fixture
def legacy(tmp_path, monkeypatch):
    """สร้างไฟล์รุ่นเก่าปลอม (ตัวตรวจชนิดถูกบังคับ) + soffice ปลอม"""

    def setup(kind: DetectedType, mode: str = "ok"):
        f = tmp_path / f"a.{kind.value}"
        f.write_bytes(b"legacy bytes")
        monkeypatch.setattr(pipeline_mod, "detect", lambda p: kind)
        monkeypatch.setattr(convert_mod, "find_soffice", lambda: "soffice-fake")
        monkeypatch.setattr(convert_mod, "check_not_encrypted", lambda p, k: None)
        run = FakeRun(kind, mode)
        monkeypatch.setattr(convert_mod, "_run", run)
        return f, run

    return setup


@pytest.mark.parametrize("kind", [DetectedType.DOC, DetectedType.XLS, DetectedType.PPT])
def test_conversion_success_goes_through_modern_extractor(legacy, kind):
    f, run = legacy(kind)
    d = extract_file(f)
    assert d.error is None and d.status == DocumentStatus.DONE
    assert d.detected_type == kind and d.converted_from == kind.value
    assert d.units and all(u.method == ExtractionMethod.CONVERTED for u in d.units)
    ext, _ = MAKERS[kind]
    assert f"--convert-to" in run.cmd and any(a.startswith(f"{ext}:") for a in run.cmd)
    assert run.timeout == convert_mod.CONVERT_TIMEOUT_SECONDS


def test_command_uses_isolated_profile_headless_no_restore(legacy):
    f, run = legacy(DetectedType.DOC)
    extract_file(f)
    assert "--headless" in run.cmd and "--norestore" in run.cmd
    profile = next(a for a in run.cmd if a.startswith("-env:UserInstallation=file:///"))
    assert "profile" in profile and "rag-convert-" in profile
    assert Path(run.cmd[-1]).parent.parent == run.work_dir  # ต้นทางที่แปลงเป็นสำเนาในโฟลเดอร์ชั่วคราว ไม่ใช่ไฟล์เดิม
    assert Path(run.cmd[-1]) != f


@pytest.mark.parametrize("mode", ["ok", "nofile", "fail", "timeout"])
def test_temp_dir_removed_after_success_and_failure(legacy, mode):
    f, run = legacy(DetectedType.DOC, mode)
    extract_file(f)
    assert run.work_dir.name.startswith("rag-convert-")
    assert not run.work_dir.exists()


def test_no_output_file_is_corrupt(legacy):
    f, _ = legacy(DetectedType.XLS, "nofile")
    d = extract_file(f)
    assert d.status == DocumentStatus.FAILED and d.error.code == ErrorCode.CORRUPT
    assert d.error.message == "แปลงไฟล์ไม่สำเร็จ ไฟล์อาจเสียหายหรือมีรหัสผ่าน"
    assert d.detected_type == DetectedType.XLS and d.converted_from == "xls"


def test_nonzero_exit_is_corrupt(legacy):
    f, _ = legacy(DetectedType.PPT, "fail")
    assert extract_file(f).error.code == ErrorCode.CORRUPT


def test_timeout_is_conversion_timeout_and_retriable(legacy):
    f, _ = legacy(DetectedType.DOC, "timeout")
    d = extract_file(f)
    assert d.error.code == ErrorCode.CONVERSION_TIMEOUT and d.error.retriable


def test_no_libreoffice_is_unsupported_type(tmp_path, monkeypatch):
    f = tmp_path / "a.doc"
    f.write_bytes(b"x")
    monkeypatch.setattr(pipeline_mod, "detect", lambda p: DetectedType.DOC)
    monkeypatch.setattr(convert_mod, "find_soffice", lambda: None)
    d = extract_file(f)
    assert d.error.code == ErrorCode.UNSUPPORTED_TYPE
    assert d.error.message == "ระบบยังไม่ได้ติดตั้งตัวแปลงไฟล์รุ่นเก่า กรุณาบันทึกเป็น .docx / .xlsx / .pptx แล้วอัปโหลดใหม่"


def test_conversion_that_yields_empty_workbook_is_empty_error(legacy, monkeypatch):
    f, run = legacy(DetectedType.XLS)

    def empty_run(cmd, timeout):
        out_dir = Path(cmd[cmd.index("--outdir") + 1])
        Workbook().save(out_dir / "input.xlsx")
        return 0

    monkeypatch.setattr(convert_mod, "_run", empty_run)
    assert extract_file(f).error.code == ErrorCode.EMPTY


# ---------- หา soffice ----------

def test_soffice_path_env_wins(tmp_path, monkeypatch):
    fake = tmp_path / "soffice.exe"
    fake.write_text("x")
    monkeypatch.setenv("SOFFICE_PATH", str(fake))
    assert convert_mod.find_soffice() == str(fake)
    monkeypatch.setenv("SOFFICE_PATH", str(tmp_path / "missing.exe"))
    assert convert_mod.find_soffice() is None


# ---------- ไฟล์ใส่รหัส ----------

class FakeOle:
    streams: dict = {}

    def __init__(self, path):
        pass

    def exists(self, name):
        return name in self.streams

    def openstream(self, name):
        import io

        return io.BytesIO(self.streams[name])

    def close(self):
        pass


def biff(*records):
    import struct

    return b"".join(struct.pack("<HH", rid, len(body)) + body for rid, body in records)


def test_xls_with_filepass_is_encrypted(monkeypatch, tmp_path):
    FakeOle.streams = {"Workbook": biff((0x0809, b"\0" * 8), (0x002F, b"\0" * 6), (0x000A, b""))}
    monkeypatch.setattr(convert_mod.olefile, "OleFileIO", FakeOle)
    with pytest.raises(convert_mod.ExtractionError) as e:
        convert_mod.check_not_encrypted(tmp_path / "x.xls", DetectedType.XLS)
    assert e.value.code == ErrorCode.ENCRYPTED


def test_xls_without_filepass_passes(monkeypatch, tmp_path):
    FakeOle.streams = {"Workbook": biff((0x0809, b"\0" * 8), (0x0042, b"\0" * 2), (0x000A, b""), (0x002F, b"\0" * 6))}
    monkeypatch.setattr(convert_mod.olefile, "OleFileIO", FakeOle)
    convert_mod.check_not_encrypted(tmp_path / "x.xls", DetectedType.XLS)  # FILEPASS หลัง EOF ของ globals ไม่นับ


@pytest.mark.parametrize("flags,encrypted", [(0x0100, True), (0x0000, False)])
def test_doc_fencrypted_bit(monkeypatch, tmp_path, flags, encrypted):
    import struct

    head = b"\xec\xa5" + b"\0" * 8 + struct.pack("<H", flags) + b"\0" * 4
    FakeOle.streams = {"WordDocument": head}
    monkeypatch.setattr(convert_mod.olefile, "OleFileIO", FakeOle)
    if encrypted:
        with pytest.raises(convert_mod.ExtractionError) as e:
            convert_mod.check_not_encrypted(tmp_path / "x.doc", DetectedType.DOC)
        assert e.value.code == ErrorCode.ENCRYPTED
    else:
        convert_mod.check_not_encrypted(tmp_path / "x.doc", DetectedType.DOC)


# ---------- ฆ่า process tree จริง ----------

def test_timeout_kills_whole_process_tree(tmp_path):
    beat = tmp_path / "beat.txt"
    child = tmp_path / "child.py"
    child.write_text(textwrap.dedent(f"""
        import time
        while True:
            open({str(beat)!r}, "a").write("x")
            time.sleep(0.1)
    """))
    parent = tmp_path / "parent.py"
    parent.write_text(textwrap.dedent(f"""
        import subprocess, sys, time
        subprocess.Popen([sys.executable, {str(child)!r}])
        time.sleep(60)
    """))
    t0 = time.time()
    with pytest.raises(subprocess.TimeoutExpired):
        convert_mod._run([sys.executable, str(parent)], timeout=2)
    assert time.time() - t0 < 30
    time.sleep(0.5)
    size = beat.stat().st_size  # ลูกเคยเขียน (มันรันอยู่ก่อนถูกฆ่า)
    assert size > 0
    time.sleep(1.0)
    assert beat.stat().st_size == size  # ลูกหยุดเขียน = ถูกฆ่าไปด้วย


# ---------- page_equivalent ของ .doc จาก SummaryInformation ----------

class FakeMeta:
    def __init__(self, pages):
        self.num_pages = pages


class FakeOleMeta:
    pages = None

    def __init__(self, path):
        pass

    def get_metadata(self):
        return FakeMeta(self.pages)

    def close(self):
        pass


@pytest.mark.parametrize("pages,expected", [(45, 45), (0, None), (None, None)])
def test_doc_page_count_from_summary_information(monkeypatch, tmp_path, pages, expected):
    FakeOleMeta.pages = pages
    monkeypatch.setattr(convert_mod.olefile, "OleFileIO", FakeOleMeta)
    assert convert_mod.legacy_page_count(tmp_path / "a.doc", DetectedType.DOC) == expected


def test_page_count_only_for_doc(monkeypatch, tmp_path):
    FakeOleMeta.pages = 9
    monkeypatch.setattr(convert_mod.olefile, "OleFileIO", FakeOleMeta)
    assert convert_mod.legacy_page_count(tmp_path / "a.xls", DetectedType.XLS) is None


def test_pipeline_uses_original_doc_page_count(legacy, monkeypatch):
    f, _ = legacy(DetectedType.DOC)
    monkeypatch.setattr(pipeline_mod, "legacy_page_count", lambda p, k: 45)
    assert extract_file(f).page_equivalent == 45
    monkeypatch.setattr(pipeline_mod, "legacy_page_count", lambda p, k: None)
    assert extract_file(f).page_equivalent != 45  # ไม่มีค่าในไฟล์เดิม -> ใช้ค่าของ docx ที่แปลงได้ตามเดิม
