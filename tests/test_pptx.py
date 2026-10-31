"""PPTX extractor (app/extraction/pptx.py) ด้วยไฟล์สังเคราะห์จาก python-pptx"""

import io
import zipfile
from types import SimpleNamespace

import pymupdf
import pytest
from pptx import Presentation
from pptx.chart.data import CategoryChartData, XyChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.util import Inches, Pt

from app.extraction.pipeline import extract_file
from app.extraction.pptx import _chart_block, reading_order
from app.extraction.tally import Tally
from app.models import BlockType, DetectedType, DocumentStatus, ErrorCode, UnitType

TITLE_ONLY, BLANK = 5, 6


def new_prs():
    return Presentation()  # 10in x 7.5in


def add_slide(prs, layout=BLANK, title=None):
    s = prs.slides.add_slide(prs.slide_layouts[layout])
    if title is not None:
        s.shapes.title.text = title
    return s


def box(slide, text, left, top, w=2.0, h=0.5):
    tb = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(w), Inches(h))
    tb.text_frame.text = text
    return tb


def png(w=8, h=8) -> bytes:
    return pymupdf.Pixmap(pymupdf.csRGB, pymupdf.IRect(0, 0, w, h)).tobytes("png")


def save(prs, tmp_path, name="a.pptx"):
    f = tmp_path / name
    prs.save(str(f))
    return f


def extract(path):
    d = extract_file(path)
    assert d.error is None, d.error
    return d


def texts(unit):
    return [b.text for b in unit.blocks]


LONG_EN = "This slide explains how multithreaded programs share memory between threads. " * 4  # > 100 letters


# ---------- ลำดับการอ่าน ----------

def test_reading_order_is_by_position_not_creation_order(tmp_path):
    prs = new_prs()
    s = add_slide(prs, TITLE_ONLY, "หัวเรื่อง")
    box(s, "ขวาบน", 6, 1.0)  # สร้างก่อน แต่อยู่ขวา
    box(s, "ซ้ายล่าง", 1, 4.0)
    box(s, "ซ้ายบน", 1, 1.1)  # top ต่างจากขวาบน 0.1in (< 5% ของ 7.5in) = แถวเดียวกัน
    box(s, "กลางล่าง", 4, 4.1)
    d = extract(save(prs, tmp_path))
    assert texts(d.units[0]) == ["หัวเรื่อง", "ซ้ายบน", "ขวาบน", "ซ้ายล่าง", "กลางล่าง"]
    assert d.units[0].blocks[0].type == BlockType.HEADING and d.units[0].blocks[0].level == 1


def test_row_band_boundary():
    def item(top, left, order):
        return SimpleNamespace(box=(left, top, 1.0, 1.0), order=order)

    H = 100.0  # band = 5
    a, b, c = item(0, 50, 0), item(5, 10, 1), item(5.01, 5, 2)
    # a เริ่มแถว; b ห่าง 5 พอดี = แถวเดียวกัน (เรียงตาม left); c ห่าง 5.01 = แถวใหม่
    assert reading_order([c, b, a], H) == [b, a, c]


def test_shapes_without_position_go_last():
    placed = SimpleNamespace(box=(0, 0, 1, 1), order=1)
    floating = SimpleNamespace(box=None, order=0)
    assert reading_order([floating, placed], 100.0) == [placed, floating]


def test_group_children_use_slide_coordinates(tmp_path):
    prs = new_prs()
    s = add_slide(prs)
    box(s, "กลางสไลด์", 1, 3.0)
    grp = s.shapes.add_group_shape()
    box_in_group = grp.shapes.add_textbox(Inches(1), Inches(0.5), Inches(2), Inches(0.5))
    box_in_group.text_frame.text = "ในกลุ่มบน"
    other = grp.shapes.add_textbox(Inches(1), Inches(1.5), Inches(2), Inches(0.5))
    other.text_frame.text = "ในกลุ่มล่าง"
    # ย้ายกลุ่มลงไป 4in และขยาย 2 เท่า: chOff/chExt ยังเป็นพิกัดเดิมของลูก
    xfrm = grp._element.grpSpPr.xfrm
    ns = {"a": "http://schemas.openxmlformats.org/drawingml/2006/main"}
    chext = xfrm.find("a:chExt", ns)
    off, ext = xfrm.find("a:off", ns), xfrm.find("a:ext", ns)
    off.set("x", "0")
    off.set("y", str(Inches(4)))
    ext.set("cx", str(int(chext.get("cx")) * 2))
    ext.set("cy", str(int(chext.get("cy")) * 2))
    chOff = xfrm.find("a:chOff", ns)
    chOff.set("x", "0")
    chOff.set("y", "0")
    d = extract(save(prs, tmp_path))
    # ลูกตัวบนอยู่ที่ 4 + 0.5*2 = 5in > 3in จึงตามหลัง "กลางสไลด์" (ถ้าใช้พิกัดดิบ 0.5in จะมาก่อน)
    assert texts(d.units[0]) == ["กลางสไลด์", "ในกลุ่มบน", "ในกลุ่มล่าง"]
    bbox = d.units[0].blocks[1].bbox
    assert bbox[1] == pytest.approx(5 * 72, abs=1)  # หน่วย point บนพิกัดสไลด์


def test_nested_group(tmp_path):
    prs = new_prs()
    s = add_slide(prs)
    outer = s.shapes.add_group_shape()
    inner = outer.shapes.add_group_shape()
    inner.shapes.add_textbox(Inches(1), Inches(1), Inches(2), Inches(0.5)).text_frame.text = "ลึกสุด"
    outer.shapes.add_textbox(Inches(1), Inches(3), Inches(2), Inches(0.5)).text_frame.text = "ชั้นนอก"
    d = extract(save(prs, tmp_path))
    assert texts(d.units[0]) == ["ลึกสุด", "ชั้นนอก"]


# ---------- ตาราง / กราฟ / notes ----------

def test_table_with_merged_cells(tmp_path):
    prs = new_prs()
    s = add_slide(prs)
    gf = s.shapes.add_table(3, 3, Inches(1), Inches(1), Inches(6), Inches(2))
    t = gf.table
    for r in range(3):
        for c in range(3):
            t.cell(r, c).text = f"ช่อง{r}{c}"
    t.cell(0, 0).merge(t.cell(0, 1))
    t.cell(1, 2).merge(t.cell(2, 2))
    d = extract(save(prs, tmp_path))
    b = d.units[0].blocks[0]
    assert b.type == BlockType.TABLE
    assert {(m.row, m.col, m.rowspan, m.colspan) for m in b.table.merged_cells} == {(0, 0, 1, 2), (1, 2, 2, 1)}
    assert len(b.table.rows) == 3 and all(len(r) == 3 for r in b.table.rows)
    lines = b.text.split("\n")
    assert lines[0].count("ช่อง00") == 1 and lines[0].count("|") == 1
    assert lines[1].startswith("ช่อง10 | ช่อง11 | ช่อง12") and lines[2] == "ช่อง20 | ช่อง21"


def test_chart_becomes_table_block_with_data(tmp_path):
    prs = new_prs()
    s = add_slide(prs, TITLE_ONLY, "ผลการทดลอง")
    cd = CategoryChartData()
    cd.categories = ["10", "20", "30"]
    cd.add_series("Process", (3633.8, 6947.8, 10501))
    cd.add_series("Thread", (87.8, 168.2, 269.8))
    gf = s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(1), Inches(2), Inches(6), Inches(4), cd)
    gf.chart.has_title = True
    gf.chart.chart_title.text_frame.text = "เวลาสร้าง"
    d = extract(save(prs, tmp_path))
    b = d.units[0].blocks[1]
    assert b.type == BlockType.TABLE
    assert b.text.split("\n") == ["กราฟ: เวลาสร้าง", "หมวด | Process | Thread", "10 | 3633.8 | 87.8", "20 | 6947.8 | 168.2", "30 | 10501 | 269.8"]
    assert b.table.header_rows == 1 and "pptx:chart_unreadable" not in d.units[0].signals


def test_unreadable_chart_does_not_crash():
    class Boom:
        has_title = False

        @property
        def plots(self):
            raise RuntimeError("ไม่มีข้อมูล")

    block, bad = _chart_block(SimpleNamespace(chart=Boom()), Tally(), None)
    assert bad is True and block is None


def test_xy_chart_does_not_crash(tmp_path):
    prs = new_prs()
    s = add_slide(prs)
    cd = XyChartData()
    series = cd.add_series("xy")
    series.add_data_point(1, 2)
    s.shapes.add_chart(XL_CHART_TYPE.XY_SCATTER, Inches(1), Inches(1), Inches(4), Inches(3), cd)
    d = extract(save(prs, tmp_path))
    assert d.units[0].status.value in ("ok", "needs_ocr")


def test_speaker_notes_are_last_block(tmp_path):
    prs = new_prs()
    s = add_slide(prs, TITLE_ONLY, "หัวเรื่อง")
    box(s, "เนื้อหา", 1, 2)
    s.notes_slide.notes_text_frame.text = "พูดถึงต้นทุน"
    d = extract(save(prs, tmp_path))
    assert texts(d.units[0]) == ["หัวเรื่อง", "เนื้อหา", "บันทึกผู้บรรยาย: พูดถึงต้นทุน"]


def test_hidden_slide_is_kept_with_signal(tmp_path):
    prs = new_prs()
    s = add_slide(prs, TITLE_ONLY, "สไลด์ที่ซ่อน")
    s._element.set("show", "0")
    add_slide(prs, TITLE_ONLY, "สไลด์ปกติ")
    d = extract(save(prs, tmp_path))
    assert d.unit_count == 2 and d.units[0].signals["pptx:hidden"] == 1.0 and "pptx:hidden" not in d.units[1].signals


def test_label_position_and_type(tmp_path):
    prs = new_prs()
    add_slide(prs, TITLE_ONLY, "บทนำ")
    s2 = add_slide(prs)
    box(s2, "ไม่มีหัวเรื่อง", 1, 1)
    d = extract(save(prs, tmp_path))
    assert [u.label for u in d.units] == ["สไลด์ 1: บทนำ", "สไลด์ 2"]
    assert [u.position for u in d.units] == [{"slide": 1}, {"slide": 2}]
    assert all(u.type == UnitType.SLIDE and u.method.value == "office_xml" for u in d.units)
    assert d.detected_type == DetectedType.PPTX and d.page_equivalent == 2


def test_footer_and_slide_number_placeholders_are_skipped(tmp_path):
    prs = new_prs()
    s = add_slide(prs, TITLE_ONLY, "หัวเรื่อง")
    box(s, "เนื้อหา", 1, 2)
    d = extract(save(prs, tmp_path))
    assert texts(d.units[0]) == ["หัวเรื่อง", "เนื้อหา"]


# ---------- ข้อความ / normalizer ----------

def test_vertical_tab_in_text_box_is_newline_not_replacement_char(tmp_path):
    prs = new_prs()
    s = add_slide(prs)
    tf = box(s, "", 1, 1).text_frame
    p = tf.paragraphs[0]
    p.add_run().text = "บรรทัดแรก"
    p.add_line_break()
    p.add_run().text = "บรรทัดสอง"
    d = extract(save(prs, tmp_path))
    assert d.units[0].text == "บรรทัดแรก\nบรรทัดสอง"
    assert "�" not in d.units[0].text and "norm_sig:fffd_remaining" not in d.units[0].signals


def test_paragraph_breaks_inside_a_box_are_kept(tmp_path):
    prs = new_prs()
    s = add_slide(prs)
    box(s, "บรรทัดหนึ่ง\nบรรทัดสอง", 1, 1)
    assert extract(save(prs, tmp_path)).units[0].text == "บรรทัดหนึ่ง\nบรรทัดสอง"


def test_normalizer_applies(tmp_path):
    prs = new_prs()
    s = add_slide(prs)
    box(s, "ธุุรกิิจ ๒๕๖๘", 1, 1)
    d = extract(save(prs, tmp_path))
    assert d.units[0].text == "ธุรกิจ 2568" and d.units[0].signals["norm_fix:dup_mark"] >= 1


# ---------- ด่านตรวจ ----------

def test_slide_with_only_a_big_picture_needs_ocr(tmp_path):
    prs = new_prs()
    s = add_slide(prs)
    s.shapes.add_picture(io.BytesIO(png()), 0, 0, prs.slide_width, prs.slide_height)
    d = extract(save(prs, tmp_path))
    u = d.units[0]
    assert u.status.value == "needs_ocr" and "scan_like" in [f.value for f in u.quality.flags]
    assert u.signals["image_coverage"] == pytest.approx(1.0) and u.signals["pptx:picture_count"] == 1.0
    assert d.status == DocumentStatus.NEEDS_OCR


def test_title_plus_big_picture_is_ok_with_image_content(tmp_path):
    prs = new_prs()
    s = add_slide(prs, TITLE_ONLY, "Overview of the system architecture")
    s.shapes.add_picture(io.BytesIO(png()), Inches(1), Inches(2), Inches(5), Inches(4))  # 20/75 = 0.267
    u = extract(save(prs, tmp_path)).units[0]
    assert u.status.value == "ok" and [f.value for f in u.quality.flags] == ["image_content"]
    assert u.signals["image_coverage"] == pytest.approx(20 / 75, abs=0.01)


def test_small_picture_is_not_image_content(tmp_path):
    prs = new_prs()
    s = add_slide(prs, TITLE_ONLY, "Overview of the system architecture")
    s.shapes.add_picture(io.BytesIO(png()), Inches(1), Inches(2), Inches(3), Inches(2))  # 6/75 = 0.08
    u = extract(save(prs, tmp_path)).units[0]
    assert u.status.value == "ok" and [f.value for f in u.quality.flags] == ["low_text"]


def test_long_english_slide_is_not_low_text(tmp_path):
    prs = new_prs()
    s = add_slide(prs, TITLE_ONLY, "Threads")
    box(s, LONG_EN, 1, 2, 8, 4)
    u = extract(save(prs, tmp_path)).units[0]
    assert u.status.value == "ok" and not u.quality.flags and u.signals["gate:letters"] >= 100


def test_empty_slide_without_picture_is_ok_empty(tmp_path):
    prs = new_prs()
    add_slide(prs)
    u = extract(save(prs, tmp_path)).units[0]
    assert u.status.value == "ok" and [f.value for f in u.quality.flags] == ["empty"]


# ---------- error ----------

def test_corrupt_pptx(tmp_path):
    f = tmp_path / "c.pptx"
    with zipfile.ZipFile(f, "w") as z:
        z.writestr("ppt/presentation.xml", "not xml")
    d = extract_file(f)
    assert d.detected_type == DetectedType.PPTX and d.error.code == ErrorCode.CORRUPT
